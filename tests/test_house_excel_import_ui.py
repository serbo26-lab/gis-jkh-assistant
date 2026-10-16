"""Clipboard-oriented house import, using an isolated database and no GIS."""

from openpyxl import load_workbook
from PySide6.QtWidgets import QDialog, QFileDialog, QMessageBox

import test_ui_workflow as workflows
from gis_jkh_assistant.app.house_import_preview import HouseImportPreview, HouseMatchDialog
from gis_jkh_assistant.database.act_requisites_store import ActRequisitesStore
from gis_jkh_assistant.gis_files.house_data import write_house_data_template

desktop = workflows.desktop
qt_app = workflows.qt_app


def _dialogs(monkeypatch, path, *, accepted=True):
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *_args: (str(path), ""))
    errors, messages, previews = [], [], []
    monkeypatch.setattr(QMessageBox, "critical", lambda *_args: errors.append(_args[2]))
    monkeypatch.setattr(QMessageBox, "warning", lambda *_args: errors.append(_args[2]))
    monkeypatch.setattr(QMessageBox, "information", lambda *_args: messages.append(_args[1]))

    def preview(dialog):
        previews.append(dialog.table.rowCount())
        return QDialog.DialogCode.Accepted if accepted else QDialog.DialogCode.Rejected

    def unexpected_match(_dialog):
        raise AssertionError("Unique addresses must import without manual matching")

    monkeypatch.setattr(HouseImportPreview, "exec", preview)
    monkeypatch.setattr(HouseMatchDialog, "exec", unexpected_match)
    return errors, messages, previews


def test_reordered_clipboard_paste_and_reimport_preserve_houses_and_signatures(desktop, tmp_path, monkeypatch):
    repo = desktop.repository
    second = repo.create_house(organization_id=desktop.house.organization_id,
                               address="г. Тестовый, ул. Учебная, д. 2")
    houses = [desktop.house, second]
    store = ActRequisitesStore(repo)
    customer = store.save_person(dict(full_name="Иванова Мария Ивановна", valid_from="2020-01-01"),
                                 house_id=desktop.house.id)
    director = store.save_person(dict(full_name="Петров Иван Иванович", position="Директор",
                                     basis_kind="Устав", valid_from="2020-01-01"),
                                 organization_id=desktop.house.organization_id)
    store.save_signature(customer, b"fictional-customer-signature", "customer.png")
    store.save_signature(director, b"fictional-director-signature", "director.png")
    path = tmp_path / "copy-paste.xlsx"
    write_house_data_template(path, houses=houses, tariffs_by_house={}, repository=repo)
    book = load_workbook(path)
    sheet = book["Тарифы"]
    # An entire rectangle is pasted in a different order, starting at column A.
    for row_number, house in enumerate(reversed(houses), 2):
        sheet.cell(row_number, 1, house.address)
        for index in range(23):
            sheet.cell(row_number, 2 + index * 2, f"08.{2004 + index}")
            sheet.cell(row_number, 3 + index * 2, str(10 + row_number + index))
    people = book["Представители домов"]
    headers = [c.value for c in people[1]]
    for row in people.iter_rows(min_row=2):
        if row[0].value == desktop.house.address:
            people.cell(row[0].row, headers.index("Квартира / помещение") + 1, "7")
    book.save(path)
    book.close()
    errors, messages, previews = _dialogs(monkeypatch, path)
    page = desktop.window.houses_page
    page._import_house_data()
    assert not errors
    assert len(previews) == 1
    for row_number, house in enumerate(reversed(houses), 2):
        history = repo.list_house_tariffs(house.id)
        assert len(history) == 23
        assert [(t.start_year, t.start_month, t.amount) for t in history] == [
            (2004 + index, 8, str(10 + row_number + index)) for index in range(23)
        ]
    updated = store.list_people(house_id=desktop.house.id)
    assert len(updated) == 1
    assert updated[0]["id"] == customer
    assert updated[0]["premises"] == "7"
    assert updated[0]["signature"] == b"fictional-customer-signature"
    directors = store.list_people(organization_id=desktop.house.organization_id)
    assert len(directors) == 1
    assert directors[0]["id"] == director
    assert directors[0]["signature"] == b"fictional-director-signature"
    after_first = list(repo.connection.iterdump())
    page._import_house_data()
    assert not errors
    assert len(previews) == 1
    assert messages[-1] == "Изменений нет"
    assert list(repo.connection.iterdump()) == after_first


def test_cancelled_visible_template_import_leaves_database_untouched(desktop, tmp_path, monkeypatch):
    path = tmp_path / "cancel.xlsx"
    write_house_data_template(path, houses=[desktop.house], tariffs_by_house={})
    book = load_workbook(path)
    book["Площади"].cell(2, 2, "500")
    book["Тарифы"].cell(2, 2, "08.2026")
    book["Тарифы"].cell(2, 3, "23,99")
    book.save(path)
    book.close()
    errors, _, previews = _dialogs(monkeypatch, path, accepted=False)
    before = list(desktop.repository.connection.iterdump())
    desktop.window.houses_page._import_house_data()
    assert not errors
    assert previews
    assert list(desktop.repository.connection.iterdump()) == before
