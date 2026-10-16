"""Recover the last tariff pasted into hidden columns of older templates."""

from datetime import datetime
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from gis_jkh_assistant.gis_files.house_data import (
    HouseDataImportError, parse_house_data_template, resolve_house_row,
)


def _legacy_book(path, last_start, last_amount, *, object_guid=None, fias_guid=None):
    book = Workbook()
    areas = book.active
    areas.title = "Площади"
    areas.append(["_house_id", "Адрес", "Площадь к начислению, м²"])
    areas.append(["local-house", "Дом А", "100"])
    sheet = book.create_sheet("Тарифы")
    sheet.append(["_house_id", "Адрес", "Действует с", "Тариф, руб./м² в месяц",
                  "_house_object_guid", "_fias_house_guid", "_gis_unique_house_number"])
    sheet.append(["local-house", "Дом А", "08.2025", "22,15",
                  object_guid or last_start, fias_guid or last_amount, "gis-house-number"])
    sheet.cell(2, 5).number_format = "mm.yyyy"
    book.save(path)
    book.close()


@pytest.mark.parametrize("last_start", ["08.2026", datetime(2026, 8, 1)])
def test_last_hidden_pair_is_imported_and_not_used_as_house_identity(tmp_path, last_start):
    path = tmp_path / "legacy-paste.xlsx"
    _legacy_book(path, last_start, "23,99")
    result = parse_house_data_template(path)
    assert [(row.start.input_label, row.amount) for row in result.tariffs] == [
        ("08.2025", "22,15"), ("08.2026", "23,99"),
    ]
    assert all(row.house_object_guid is None and row.fias_house_guid is None for row in result.tariffs)
    assert all(row.gis_unique_house_number == "gis-house-number" for row in result.tariffs)
    house = SimpleNamespace(id="local-house", address="Дом А", house_object_guid="real-object",
                            fias_house_guid="real-fias", gis_unique_house_number="gis-house-number")
    assert all(resolve_house_row(row, [house]) is house for row in result.tariffs)
    assert len(result.warnings) == 1
    assert "Строк: 1" in result.warnings[0]


@pytest.mark.parametrize("start,amount", [
    ("08.2026", None), (None, "23,99"), ("08.2026", "0"),
    ("08.2026", "-1"), ("08.2026", "не число"),
])
def test_incomplete_or_invalid_hidden_pair_never_loses_a_tariff_silently(tmp_path, start, amount):
    path = tmp_path / "invalid-legacy-paste.xlsx"
    _legacy_book(path, start, amount)
    with pytest.raises(HouseDataImportError, match="строка 2, пара 2"):
        parse_house_data_template(path)


def test_legitimate_conflicting_guids_still_require_explicit_house_match(tmp_path):
    path = tmp_path / "conflicting-legacy.xlsx"
    _legacy_book(path, None, None, object_guid="foreign-object", fias_guid="foreign-fias")
    result = parse_house_data_template(path)
    assert len(result.tariffs) == 1
    assert not result.warnings
    house = SimpleNamespace(id="local-house", address="Дом А", house_object_guid="real-object",
                            fias_house_guid="real-fias", gis_unique_house_number="gis-house-number")
    with pytest.raises(ValueError, match="противоречат"):
        resolve_house_row(result.tariffs[0], [house])


def test_unrecognized_columns_after_old_metadata_are_not_silently_skipped(tmp_path):
    path = tmp_path / "extended-old.xlsx"
    _legacy_book(path, None, None, object_guid="object", fias_guid="fias")
    from openpyxl import load_workbook
    book = load_workbook(path)
    book["Тарифы"].cell(2, 8, "08.2027")
    book["Тарифы"].cell(2, 9, "25")
    book.save(path)
    book.close()
    with pytest.raises(HouseDataImportError, match="после служебных столбцов"):
        parse_house_data_template(path)
