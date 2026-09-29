from openpyxl import Workbook

from gis_jkh_assistant.gis_files.work_lists import parse_work_template


def _build_workbook(path, filled: bool) -> None:
    wb = Workbook()
    main = wb.active
    main.title = "Перечень работ и услуг"
    main.append(
        ["Работа (услуга)", "Цена, руб.", "Объем", "Кол-во", "Итого-стоимость, руб.", "Статус обработки"]
    )
    if filled:
        main.append(["12.Услуга по управлению", 10.5, 2, 1, 21, None])

    options = wb.create_sheet("ОпцииПеречня")
    options.append(["houseGuid", "house-guid", "Ссылка на дом"])
    options.append(["yearFrom", 2020, 'Период "с"'])
    options.append(["monthFrom", 5, None])
    options.append(["yearFrom", 2020, 'Период "по"'])
    options.append(["monthFrom", 12, None])

    lookup = wb.create_sheet("СпрРабУсл")
    lookup.append(["12.Услуга по управлению", "work-guid-12"])
    lookup.append(["16.Уборка лестничных клеток", "work-guid-16"])

    conf = wb.create_sheet("conf")
    conf.append(["TYPE", "PFRPERv2"])
    conf.append(["VERSION", "99.0.0"])
    wb.save(path)


def test_blank_gis_file_uses_embedded_work_dictionary(tmp_path) -> None:
    path = tmp_path / "blank.xlsx"
    _build_workbook(path, filled=False)

    result = parse_work_template(path)

    assert result.house_guid == "house-guid"
    assert result.start_year == 2020
    assert result.start_month == 5
    assert result.end_month == 12
    assert len(result.items) == 2
    assert result.items[0].work_code == "12"
    assert result.items[0].work_guid == "work-guid-12"


def test_filled_gis_file_imports_selected_rows_and_values(tmp_path) -> None:
    path = tmp_path / "filled.xlsx"
    _build_workbook(path, filled=True)

    result = parse_work_template(path)

    assert len(result.items) == 1
    item = result.items[0]
    assert item.work_name == "Услуга по управлению"
    assert item.price == "10.5"
    assert item.volume == "2"
    assert item.quantity == "1"
    assert item.cost == "21"
