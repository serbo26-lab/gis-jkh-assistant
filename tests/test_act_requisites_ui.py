"""Real Qt dialogs and workbook apply flow on isolated fictional databases."""
import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QDialog, QFileDialog, QMessageBox, QPushButton

import test_ui_workflow as workflows
from act_test_data import populate_act_requisites
from gis_jkh_assistant.app.house_import_preview import HouseImportPreview
from gis_jkh_assistant.app.quality_bulk_dialogs import GenerateAcceptanceActDialog
from gis_jkh_assistant.database.act_requisites_store import ActRequisitesStore
from gis_jkh_assistant.gis_files.house_data import write_house_data_template

desktop = workflows.desktop
qt_app = workflows.qt_app


@pytest.mark.parametrize("fail", [False, True])
def test_extended_import_applies_as_one_transaction_and_keeps_existing_acts(desktop, tmp_path, monkeypatch, fail):
    repo, window = desktop.repository, desktop.window
    populate_act_requisites(repo, desktop.work_list)
    store = window.quality_page.store
    period = store.get_or_create_period(work_list_id=desktop.work_list.id, year=2026, month=1)
    item = store.list_items(period.id)[0]
    act = store.add_act(item_id=item.id, act_type="ACCEPTANCE_ACT", number="TEST", reg_date="31.01.2026",
        period_from="01.01.2026", period_to="31.01.2026", file_name="user.pdf", description=None,
        content_type="application/pdf", content=b"USER ACT")
    file = tmp_path / "full.xlsx"
    house = repo.get_house(desktop.house.id)
    write_house_data_template(file, houses=[house], tariffs_by_house={}, repository=repo)
    book = load_workbook(file)
    houses = book["Дома"]
    headers = {cell.value: cell.column for cell in houses[1]}
    houses.cell(2, headers["Количество жилых помещений"], 17)
    houses.cell(2, headers["Количество нежилых помещений"], 0)
    houses.cell(2, headers["Номер договора управления"], "NEW TEST")
    people = book["Подписанты УО"]
    headers = {cell.value: cell.column for cell in people[1]}
    people.cell(2, headers["Должность"], "Тестовый директор")
    book.save(file)
    book.close()
    if fail:
        repo.connection.execute("CREATE TRIGGER fail_person BEFORE UPDATE ON act_people BEGIN SELECT RAISE(ABORT, 'TEST PERSON FAILURE'); END")
        repo.connection.commit()
    before = list(repo.connection.iterdump())
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *_: (str(file), ""))
    previews = []
    def accept(dialog):
        previews.append(dialog.table.rowCount())
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(HouseImportPreview, "exec", accept)
    notifications = []
    for kind in ("information", "warning", "critical"):
        monkeypatch.setattr(QMessageBox, kind, lambda *a, kind=kind: notifications.append((kind, a[2])))
    window.houses_page._import_house_data()
    assert previews and previews[0] >= 3
    if fail:
        assert list(repo.connection.iterdump()) == before
        assert any(kind == "critical" for kind, _ in notifications)
    else:
        updated = repo.get_house(house.id)
        assert updated.residential_premises_count == 17
        assert updated.non_residential_premises_count == 0
        assert updated.management_agreement_number == "NEW TEST"
        assert ActRequisitesStore(repo).list_people(organization_id=house.organization_id)[0]["position"] == "Тестовый директор"
    assert store.get_act(act.id).content == b"USER ACT"
    assert store.act_link_item_ids(act.id) == [item.id]


def test_generation_dialog_resolves_parties_and_preview_does_not_write(desktop):
    repo = desktop.repository
    populate_act_requisites(repo, desktop.work_list)
    store = desktop.window.quality_page.store
    period = store.get_or_create_period(work_list_id=desktop.work_list.id, year=2026, month=1)
    calls = []
    dialog = GenerateAcceptanceActDialog(items=store.list_items(period.id), year=2026, month=1,
        repository=repo, house_id=desktop.house.id, preview_callback=calls.append)
    dialog.show()
    workflows._settle(desktop.app)
    assert dialog.executor_combo.currentData()
    assert dialog.customer_combo.currentData()
    assert "Стороны и договор заполнены" in dialog.party_status.text()
    assert not dialog._details()["issues"]
    before = list(repo.connection.iterdump())
    buttons = dialog.findChildren(QPushButton)
    next(button for button in buttons if button.text() == "Предпросмотр PDF…").click()
    assert len(calls) == 1
    assert list(repo.connection.iterdump()) == before
    dialog.close()
    dialog.deleteLater()


def test_chief_import_requires_confirmation_and_keeps_dialog_alive(desktop, monkeypatch):
    from dataclasses import replace
    from PySide6.QtWidgets import QInputDialog
    from gis_jkh_assistant.app import act_requisites_dialog as module
    page = desktop.window.houses_page
    page.refresh_organizations(select_id=desktop.organization.id)
    status = replace(desktop.gateway.status(), helper_version="0.1.32")
    monkeypatch.setattr(desktop.gateway, "status", lambda: status)
    monkeypatch.setattr(module, "GisIntegrationWorker", workflows.ControlledWorker)
    monkeypatch.setattr(QMessageBox, "information", lambda *_: None)
    dialog = module.ActRequisitesDialog(desktop.repository, organization_id=desktop.organization.id, page=page)
    dialog.show()
    workflows._settle(desktop.app)
    dialog._load_chief()
    worker = dialog._worker
    assert worker.command_type == "get_organization_act_signatory"
    assert worker.isRunning()
    dialog.reject()
    assert dialog.isVisible()
    chosen = []
    monkeypatch.setattr(QInputDialog, "getItem", lambda *args: (args[3][1], True))
    monkeypatch.setattr(dialog, "_edit", lambda **kwargs: chosen.append(kwargs["candidate"]))
    before = ActRequisitesStore(desktop.repository).list_people(organization_id=desktop.organization.id)
    worker.succeeded.emit({"ogrn": desktop.organization.ogrn, "candidates": [
        {"fullName": "Иванов Иван Иванович", "source": "Основные сведения"},
        {"fullName": "Петров Пётр Петрович", "position": "Директор", "source": "Дополнительные сведения"},
    ]})
    assert chosen == [{"full_name": "Петров Пётр Петрович", "position": "Директор"}]
    assert ActRequisitesStore(desktop.repository).list_people(organization_id=desktop.organization.id) == before
    worker.finish()
    assert dialog.load_button.isEnabled()
    dialog.close()
    dialog.deleteLater()


def test_year_generation_skips_unready_month_without_erasing_old_act(desktop, monkeypatch):
    populate_act_requisites(desktop.repository, desktop.work_list)
    page = workflows._select_page(desktop, 4)
    january, item = workflows._include_quality(page, 1)
    old = page.store.add_act(item_id=item.id, act_type="ACCEPTANCE_ACT", number="1", reg_date="31.01.2026",
        period_from="01.01.2026", period_to="31.01.2026", file_name="old.pdf", description=None,
        content_type="application/pdf", content=b"OLD AUTO", generated=True)
    # January has no total cost. February does and can be generated.
    february, _ = workflows._include_quality(page, 2)
    fixation = desktop.repository.get_or_create_work_fixation_period(work_list_id=desktop.work_list.id, year=2026, month=2)
    desktop.repository.update_work_fixation_item(desktop.repository.list_work_fixation_items(fixation.id)[0].id,
        fact_price="100", fact_amount="1", fact_count=1, fact_total_cost="100")
    january_fixation = desktop.repository.get_work_fixation_period_for_month(desktop.work_list.id, year=2026, month=1)
    desktop.repository.update_work_fixation_item(desktop.repository.list_work_fixation_items(january_fixation.id)[0].id,
        fact_price=None, fact_amount=None, fact_count=None, fact_total_cost=None)
    notices = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: notices.append(args[2]))
    page._auto_generate_year_acts(units_ready=True)
    assert page.store.get_act(old.id).content == b"OLD AUTO"
    assert not page.store.get_act(old.id).deleted
    assert len(page.store.list_period_acts(february.id)) == 1
    assert "Сформировано актов: 1" in notices[0]
    assert "Не готовы: 1" in notices[0] and "Январь" in notices[0]


def test_pdf_preview_loads_document_and_closes_without_writing(desktop):
    from gis_jkh_assistant.acceptance_act_pdf import AcceptanceActWorkRow, build_acceptance_act_pdf
    from gis_jkh_assistant.app.pdf_preview_dialog import PdfPreviewDialog
    content = build_acceptance_act_pdf(organization_name="Тестовая УО", house_address="Тестовый дом",
        act_number="1", act_date="31.01.2026", period_from="01.01.2026", period_to="31.01.2026",
        work_rows=[AcceptanceActWorkRow(name="Тестовая работа", actual_cost="100")])
    before = list(desktop.repository.connection.iterdump())
    dialog = PdfPreviewDialog(content)
    dialog.show()
    workflows._settle(desktop.app)
    assert dialog.document.pageCount() == 1
    dialog.close()
    dialog.deleteLater()
    workflows._settle(desktop.app)
    assert list(desktop.repository.connection.iterdump()) == before


def test_manual_pdf_generation_attaches_document_links_and_requisites_together(desktop, monkeypatch):
    from gis_jkh_assistant.app.quality_bulk_dialogs import QualityMonthActsDialog
    populate_act_requisites(desktop.repository, desktop.work_list)
    page = workflows._select_page(desktop, 4)
    period, item = workflows._include_quality(page, 1)
    fixation = desktop.repository.get_work_fixation_period_for_month(desktop.work_list.id, year=2026, month=1)
    desktop.repository.update_work_fixation_item(desktop.repository.list_work_fixation_items(fixation.id)[0].id,
        fact_price="100", fact_amount="2", fact_count=1, fact_total_cost="200")
    monkeypatch.setattr(GenerateAcceptanceActDialog, "exec", lambda _: QDialog.DialogCode.Accepted)
    warnings = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *args: warnings.append(args[2]))
    dialog = QualityMonthActsDialog(page.store, period_id=period.id, year=2026, month=1,
        organization_name=desktop.organization.name, house_address=desktop.house.address, parent=page)
    dialog._generate(units_ready=True)
    assert not warnings
    acts = page.store.list_period_acts(period.id)
    assert len(acts) == 1 and acts[0].content.startswith(b"%PDF")
    assert page.store.act_link_item_ids(acts[0].id) == [item.id]
    snapshot = ActRequisitesStore(desktop.repository).snapshot(acts[0].id)
    assert snapshot["executor"]["full_name"] == "Иванов Иван Иванович"
    assert snapshot["agreement_number"] == "TEST-1"
    assert snapshot["work_rows"][0]["actual_cost"] == "200"
    dialog.close()
    dialog.deleteLater()
