"""Regression evidence for the user's first-run walkthrough, without live GIS."""
from types import SimpleNamespace
from collections import Counter

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt, QTimer
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QDialog, QDialogButtonBox, QFileDialog, QMessageBox

from gis_jkh_assistant.app.russian_buttons import install
from gis_jkh_assistant.app.styles import apply_theme
from gis_jkh_assistant.app.widgets import SafeSpinBox
from gis_jkh_assistant.database import Repository
from gis_jkh_assistant.gis_files.house_data import (
    parse_house_data_template, resolve_house_row, write_house_data_template,
    area_values_for_house, ImportedAreaRow,
)
from gis_jkh_assistant.gis_integration.local_gateway import _GatewayState
import test_ui_workflow as workflow_fixtures
from test_ui_workflow import _select_page

desktop = workflow_fixtures.desktop
qt_app = workflow_fixtures.qt_app


def settle(app):
    for _ in range(6):
        app.processEvents()


@pytest.mark.parametrize("theme", ["light", "dark"])
@pytest.mark.parametrize("scale", [1.0])
def test_standard_buttons_and_year_text_are_readable(qt_app, theme, scale):
    install(qt_app)
    apply_theme(qt_app, theme, scale)
    buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Yes | QDialogButtonBox.Cancel)
    assert all(button.text() for button in buttons.buttons())
    year = SafeSpinBox()
    year.setRange(2015, 2100)
    year.setValue(2026)
    year.resize(year.minimumWidth(), 38)
    year.show()
    settle(qt_app)
    assert year.text() == "2026"
    assert year.lineEdit().width() - year.lineEdit().textMargins().right() >= year.fontMetrics().horizontalAdvance("2026")
    year.lineEdit().setFocus()
    year.lineEdit().selectAll()
    QTest.keyClicks(year.lineEdit(), "2027")
    year.interpretText()
    assert year.value() == 2027
    year._plus_button.click()
    assert year.value() == 2028
    year.close()
    year.deleteLater()


def test_export_keeps_sorted_house_id_address_and_area_together(tmp_path):
    repo = Repository(tmp_path / "source.db")
    org = repo.create_organization(name="Вымышленная УО", ogrn="1022600000001")
    first = repo.create_house(organization_id=org.id, address="Дом Б", charge_area="23", house_object_guid="B")
    second = repo.create_house(organization_id=org.id, address="Дом А", charge_area="47", house_object_guid="A")
    path = tmp_path / "data.xlsx"
    write_house_data_template(path, houses=[first, second], tariffs_by_house={})
    result = parse_house_data_template(path)
    assert [(r.house_id, r.address, r.area) for r in result.areas] == [(second.id, "Дом А", "47"), (first.id, "Дом Б", "23")]
    recreated = [SimpleNamespace(**{**vars_like(second), "id": "new-a"}), SimpleNamespace(**{**vars_like(first), "id": "new-b"})]
    assert [resolve_house_row(r, recreated).id for r in result.areas] == ["new-a", "new-b"]
    workbook = load_workbook(path)
    areas = workbook["Площади"]
    areas["C2"] = None
    areas["D2"] = "100"
    areas["E2"] = "20"
    areas["C3"] = "0"
    areas["D3"] = "500"
    workbook.save(path)
    workbook.close()
    result = parse_house_data_template(path)
    assert result.areas[0].area == "120"
    assert result.areas[0].residential_area == "100"
    assert result.areas[0].non_residential_area == "20"
    assert result.areas[1].area == "0"
    wrong_address = SimpleNamespace(**{**vars_like(first), "address": "Чужой адрес"})
    with pytest.raises(ValueError, match="адрес"):
        resolve_house_row(result.areas[1], [wrong_address])
    repo.close()


def vars_like(house):
    from dataclasses import asdict
    return asdict(house)


def test_optional_area_columns_upgrade_preserves_existing_house(tmp_path):
    path = tmp_path / "old.db"
    repo = Repository(path)
    org = repo.create_organization(name="УО", ogrn="1022600000001")
    house = repo.create_house(organization_id=org.id, address="Дом", residential_premises_count=12)
    # Model a previous schema; the next open adds only nullable fields.
    repo.connection.execute("ALTER TABLE houses DROP COLUMN residential_area")
    repo.connection.execute("ALTER TABLE houses DROP COLUMN non_residential_area")
    repo.connection.commit()
    repo.close()
    repo = Repository(path)
    restored = repo.get_house(house.id)
    assert restored.residential_premises_count == 12
    assert restored.residential_area is None
    repo.update_house(house.id, residential_area="100", non_residential_area="20")
    assert repo.get_house(house.id).non_residential_area == "20"
    repo.close()


def test_partial_area_import_keeps_unsupplied_component_when_deriving_total():
    house = SimpleNamespace(residential_area="100", non_residential_area="30")
    row = ImportedAreaRow(None, "Дом", "20", 2, non_residential_area="20", charge_explicit=False)
    assert area_values_for_house(row, house) == {"charge_area": "120", "non_residential_area": "20"}


def test_cached_directory_adds_work_on_first_click_without_network(desktop, monkeypatch):
    import json
    from gis_jkh_assistant.app.work_directory_dialog import WorkDirectorySelectionDialog
    page = _select_page(desktop, 1)
    desktop.repository.update_work_list_gis_state(desktop.work_list.id, gis_status="DRAFT")
    options = [{"guid": "fictional-new-guid", "code": "99", "name": "Вымышленная новая работа"}]
    desktop.repository.set_setting(f"ui.directory.{desktop.organization.id}", json.dumps(options))
    desktop.gateway.ready = False
    page._refresh_gis_action_buttons()
    assert page.add_works_button.isEnabled()
    monkeypatch.setattr(WorkDirectorySelectionDialog, "exec", lambda _self: QDialog.Accepted)
    monkeypatch.setattr(WorkDirectorySelectionDialog, "selected_options", lambda _self: options)
    before = len(desktop.repository.list_work_list_items(desktop.work_list.id))
    page.add_works_button.click()
    assert len(desktop.repository.list_work_list_items(desktop.work_list.id)) == before + 1
    assert page._integration_worker is None


def test_visible_selector_commands_months_and_unplanned(desktop):
    for index in (1, 2, 3, 4):
        _select_page(desktop, index)
        page = desktop.window.pages.widget(index + 1)
        settle(desktop.app)
        assert page.list_selector.isVisible()
        assert page.list_selector.item(0, 0).text()
        page.house_combo.setFocus()
        QTest.keyClick(page.house_combo, Qt.Key_Tab)
        assert desktop.app.focusWidget() is page.list_selector
        if index in (3, 4):
            assert page.months_card.isVisible()
            assert page.load_period_button.isVisible()
            assert page.current_commands_label.isVisible()
            assert not page.data_actions.isVisible()
    fixation = desktop.window.fixation_page
    _select_page(desktop, 3)
    assert fixation.unplanned_card.isVisible()
    assert fixation.load_month_button.isVisible()


def test_idle_table_geometry_converges(desktop):
    _select_page(desktop, 3)
    QTest.qWait(250)
    ticks = Counter()
    tables = {"plan": desktop.window.planning_page.table, "lists": desktop.window.work_lists_page.items_table,
              "fact": desktop.window.fixation_page.table, "quality": desktop.window.quality_page.table,
              "unplanned": desktop.window.fixation_page.unplanned_table}
    for timer in desktop.window.findChildren(QTimer):
        parent = timer.parent()
        key = type(parent).__name__
        if hasattr(parent, "widget"):
            key += ":" + next((name for name, widget in tables.items() if widget is parent.widget), "other")
        timer.timeout.connect(lambda name=key: ticks.update([name]))
    QTest.qWait(100)
    assert sum(ticks.values()) <= 2, f"Sizing must converge instead of rescheduling itself at idle: {ticks}"


def test_bulk_range_does_not_change_displayed_month_or_outside_data(desktop, monkeypatch):
    page = desktop.window.fixation_page
    _select_page(desktop, 3)
    original_month = page._current_month
    january = desktop.repository.get_work_fixation_period(page._current_period_id)
    before = desktop.repository.list_work_fixation_items(january.id)
    page.period_from.setCurrentIndex(1)
    page.period_to.setCurrentIndex(2)
    monkeypatch.setattr(QMessageBox, "question", lambda *_args, **_kwargs: QMessageBox.Yes)
    monkeypatch.setattr(QMessageBox, "information", lambda *_args, **_kwargs: QMessageBox.Ok)
    page._fill_year_from_plan()
    assert page._current_month == original_month
    assert desktop.repository.list_work_fixation_items(january.id) == before
    periods = desktop.repository.list_work_fixation_periods(desktop.work_list.id, year=2026)
    filled = {p.month for p in periods if page._period_has_data(p.id)}
    assert {1, 2, 3}.issubset(filled)
    assert 4 not in filled


def test_import_failure_rolls_back_all_area_and_tariff_updates(desktop, tmp_path, monkeypatch):
    from gis_jkh_assistant.app.house_import_preview import HouseImportPreview
    page = desktop.window.houses_page
    path = tmp_path / "import.xlsx"
    write_house_data_template(path, houses=[desktop.house], tariffs_by_house={})
    workbook = load_workbook(path)
    workbook["Площади"]["C2"] = "500"
    workbook["Площади"]["D2"] = "450"
    workbook["Тарифы"]["C2"] = "01.2026"
    workbook["Тарифы"]["D2"] = "10"
    workbook.save(path)
    workbook.close()
    desktop.repository.connection.execute("CREATE TRIGGER fail_import BEFORE INSERT ON house_tariffs BEGIN SELECT RAISE(ABORT, 'simulated import failure'); END")
    desktop.repository.connection.commit()
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *_args: (str(path), ""))
    monkeypatch.setattr(HouseImportPreview, "exec", lambda _self: QDialog.Accepted)
    errors = []
    monkeypatch.setattr(QMessageBox, "critical", lambda *args: errors.append(args[2]))
    page._import_house_data()
    assert errors
    house = desktop.repository.get_house(desktop.house.id)
    assert house.charge_area == desktop.house.charge_area
    assert house.residential_area is None
    assert not desktop.repository.list_house_tariffs(house.id)


def test_inactive_browser_cannot_reset_connected_status_or_take_command():
    state = _GatewayState()
    state.update_hello({"helperInstance": "chrome", "gisPageReady": True, "helperVersion": "0.1.31"})
    state.update_hello({"helperInstance": "edge", "gisPageReady": False, "helperVersion": "0.1.31"})
    assert state.status(True, None).gis_page_ready
    assert state.can_poll("chrome")
    assert not state.can_poll("edge")
    assert not state.can_poll("unknown")
    assert not state.can_poll(None)
    state.update_hello({"helperInstance": "chrome", "gisPageReady": False, "helperVersion": "0.1.31"})
    assert not state.status(True, None).gis_page_ready
