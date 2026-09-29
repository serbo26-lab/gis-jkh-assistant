from dataclasses import dataclass

import pytest
from openpyxl import Workbook, load_workbook

from gis_jkh_assistant.gis_files.work_lists import (
    WorkListExportError,
    export_work_list_xlsx,
)
from gis_jkh_assistant.gis_files.work_lists.exporter import (
    _patch_main_sheet,
    _patch_options_sheet,
    _replace_shared_string,
)


@dataclass
class Item:
    work_code: str
    work_name: str
    work_guid: str
    price: str | None = None
    volume: str | None = None
    quantity: str | None = None
    cost: str | None = None


def _source_template(path) -> None:
    wb = Workbook()
    main = wb.active
    main.title = "Перечень работ и услуг"
    main.append(
        [
            "Работа (услуга)",
            "Цена, руб.",
            "Объем",
            "Кол-во",
            "Итого-стоимость, руб.",
            "Статус обработки",
        ]
    )

    options = wb.create_sheet("ОпцииПеречня")
    options.append(["houseGuid", "old-house-guid", "Ссылка на дом"])
    options.append(["yearFrom", 2020, 'Период "с"'])
    options.append(["monthFrom", 1, None])
    options.append(["yearFrom", 2020, 'Период "по"'])
    options.append(["monthFrom", 12, None])

    lookup = wb.create_sheet("СпрРабУсл")
    lookup.append(["12.Услуга по управлению", "guid-12"])
    lookup.append(["13.Техническое обслуживание", "guid-13"])

    conf = wb.create_sheet("conf")
    conf.append(["TYPE", "PFRPERv2"])
    conf.append(["VERSION", "11.5.0.2"])
    wb.save(path)


def test_export_patches_house_period_and_only_filled_rows(tmp_path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "result.xlsx"
    _source_template(source)

    result = export_work_list_xlsx(
        source_template_path=source,
        output_path=output,
        house_guid="new-house-guid",
        start_year=2026,
        start_month=5,
        end_year=2026,
        end_month=12,
        items=[
            Item(
                "12",
                "Услуга по управлению",
                "guid-12",
                quantity="1",
                cost="1250,50",
            ),
            Item(
                "13",
                "Техническое обслуживание",
                "guid-13",
                price="10",
                volume="2",
                quantity="3",
            ),
            Item("13", "Техническое обслуживание", "guid-13"),
        ],
    )

    assert result.exported_rows == 2

    wb = load_workbook(output, data_only=False)
    options = wb["ОпцииПеречня"]
    assert options["B1"].value == "new-house-guid"
    assert options["B2"].value == 2026
    assert options["B3"].value == 5
    assert options["B4"].value == 2026
    assert options["B5"].value == 12

    main = wb["Перечень работ и услуг"]
    assert main["A2"].value == "12.Услуга по управлению"
    assert main["D2"].value == 1
    assert main["E2"].value == 1250.50
    assert main["B2"].value is None
    assert main["C2"].value is None

    assert main["A3"].value == "13.Техническое обслуживание"
    assert main["B3"].value == 10
    assert main["C3"].value == 2
    assert main["D3"].value == 3
    assert main["E3"].value is None
    wb.close()


def test_export_rejects_quantity_without_a_complete_mode(tmp_path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "result.xlsx"
    _source_template(source)

    with pytest.raises(WorkListExportError, match="одного количества недостаточно"):
        export_work_list_xlsx(
            source_template_path=source,
            output_path=output,
            house_guid="house-guid",
            start_year=2026,
            start_month=1,
            end_year=2026,
            end_month=12,
            items=[
                Item(
                    "12",
                    "Услуга по управлению",
                    "guid-12",
                    quantity="1",
                )
            ],
        )


def test_official_shared_string_representation_is_preserved() -> None:
    from xml.etree import ElementTree as ET

    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    worksheet = (
        f'<worksheet xmlns="{main_ns}">'
        '<dimension ref="A1:F1"/>'
        '<sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData>'
        '</worksheet>'
    ).encode()

    patched_main = _patch_main_sheet(
        worksheet,
        [("12.Услуга по управлению", "", "", "1", "1")],
        {"12.Услуга по управлению": 18},
    )
    root = ET.fromstring(patched_main)
    a2 = root.find(f".//{{{main_ns}}}c[@r='A2']")
    assert a2 is not None
    assert a2.attrib["t"] == "s"
    assert a2.find(f"{{{main_ns}}}v").text == "18"

    shared_root = ET.fromstring(
        (
            f'<sst xmlns="{main_ns}" count="2" uniqueCount="2">'
            "<si><t>header</t></si>"
            "<si><t>old-house-guid</t></si>"
            "</sst>"
        ).encode()
    )
    _replace_shared_string(shared_root, 1, "new-house-guid")
    values = [
        "".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t"))
        for item in shared_root.findall(f"{{{main_ns}}}si")
    ]
    assert values == ["header", "new-house-guid"]

    options = (
        f'<worksheet xmlns="{main_ns}"><sheetData>'
        '<row r="1"><c r="B1" t="s"><v>1</v></c></row>'
        '<row r="2"><c r="B2" t="n"><v>2020</v></c></row>'
        '<row r="3"><c r="B3" t="n"><v>1</v></c></row>'
        '<row r="4"><c r="B4" t="n"><v>2020</v></c></row>'
        '<row r="5"><c r="B5" t="n"><v>12</v></c></row>'
        "</sheetData></worksheet>"
    ).encode()
    patched_options = _patch_options_sheet(
        options,
        house_guid=None,
        start_year=2026,
        start_month=5,
        end_year=2026,
        end_month=12,
    )
    root = ET.fromstring(patched_options)
    b1 = root.find(f".//{{{main_ns}}}c[@r='B1']")
    assert b1 is not None
    assert b1.attrib["t"] == "s"
    assert b1.find(f"{{{main_ns}}}v").text == "1"
