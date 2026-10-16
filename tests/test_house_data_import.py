from datetime import datetime

from openpyxl import load_workbook

from gis_jkh_assistant.database import Repository
from gis_jkh_assistant.gis_files.house_data import (
    parse_house_data_template,
    write_house_data_template,
)
from gis_jkh_assistant.period_dates import parse_period_date


def _column(sheet, header: str, occurrence: int = 0) -> int:
    return [cell.column for cell in sheet[1] if cell.value == header][occurrence]


def test_period_date_precision() -> None:
    month = parse_period_date("08.2025", boundary="start")
    month_end = parse_period_date("07.2026", boundary="end")
    exact = parse_period_date("10.08.2025", boundary="start")
    exact_end = parse_period_date("15.11.2026", boundary="end")

    assert month.value.isoformat() == "2025-08-01"
    assert month.precision == "month"
    assert month_end.value.isoformat() == "2026-07-31"
    assert month_end.precision == "month"
    assert exact.value.isoformat() == "2025-08-10"
    assert exact.precision == "day"
    assert exact_end.value.isoformat() == "2026-11-15"
    assert exact_end.precision == "day"


def test_house_data_template_roundtrip(tmp_path) -> None:
    repository = Repository(tmp_path / "app.db")
    organization = repository.create_organization(
        name="УК Пример",
        ogrn="1022600000099",
    )
    first = repository.create_house(
        organization_id=organization.id,
        address="г Ставрополь, ул Братская, д. 6",
    )
    second = repository.create_house(
        organization_id=organization.id,
        address="г Ставрополь, ул Братская, д. 7",
    )

    path = tmp_path / "house-data.xlsx"
    write_house_data_template(
        path,
        houses=[first, second],
        tariffs_by_house={first.id: [], second.id: []},
    )

    workbook = load_workbook(path)
    areas = workbook["Площади"]
    tariffs = workbook["Тарифы"]
    areas.cell(2, _column(areas, "Площадь к начислению, м²"), "3 956,40")
    tariffs.cell(2, _column(tariffs, "Действует с"), "08.2025")
    tariffs.cell(2, _column(tariffs, "Тариф, руб./м² в месяц"), "22,00")
    tariffs.cell(2, _column(tariffs, "Действует с", 1), "10.08.2026")
    tariffs.cell(2, _column(tariffs, "Тариф, руб./м² в месяц", 1), "24.50")
    workbook.save(path)
    workbook.close()

    result = parse_house_data_template(path)

    assert result.areas[0].house_id is None
    assert result.areas[0].address == first.address
    assert result.areas[0].area == "3956,4"
    assert result.tariffs[0].start.value.isoformat() == "2025-08-01"
    assert result.tariffs[0].start.precision == "month"
    assert result.tariffs[0].amount == "22"
    assert result.tariffs[1].start.value.isoformat() == "2026-08-10"
    assert result.tariffs[1].start.precision == "day"
    assert result.tariffs[1].amount == "24,5"

    repository.close()


def test_export_keeps_existing_tariff_precision(tmp_path) -> None:
    repository = Repository(tmp_path / "app.db")
    organization = repository.create_organization(
        name="УК Пример",
        ogrn="1022600000100",
    )
    house = repository.create_house(
        organization_id=organization.id,
        address="г Ставрополь, ул Братская, д. 8",
        charge_area="4100,5",
    )
    repository.create_house_tariff(
        house_id=house.id,
        amount="23",
        start_year=2025,
        start_month=8,
        start_day=1,
        start_precision="month",
    )
    repository.create_house_tariff(
        house_id=house.id,
        amount="24",
        start_year=2026,
        start_month=8,
        start_day=10,
        start_precision="day",
    )

    path = tmp_path / "house-data.xlsx"
    write_house_data_template(
        path,
        houses=[house],
        tariffs_by_house={house.id: repository.list_house_tariffs(house.id)},
    )
    result = parse_house_data_template(path)

    assert [row.start.input_label for row in result.tariffs] == [
        "08.2025",
        "10.08.2026",
    ]

    repository.close()


def test_wide_tariff_rows_allow_gaps_and_excel_month_dates(tmp_path) -> None:
    repository = Repository(tmp_path / "app.db")
    organization = repository.create_organization(
        name="УК Пример",
        ogrn="1022600000101",
    )
    house = repository.create_house(
        organization_id=organization.id,
        address="г Ставрополь, ул Братская, д. 9",
    )

    path = tmp_path / "wide-house-data.xlsx"
    write_house_data_template(
        path,
        houses=[house],
        tariffs_by_house={house.id: []},
    )

    workbook = load_workbook(path)
    tariffs = workbook["Тарифы"]

    start = tariffs.cell(2, _column(tariffs, "Действует с"), datetime(2020, 8, 1))
    start.number_format = r"mm\.yyyy"
    tariffs.cell(2, _column(tariffs, "Тариф, руб./м² в месяц"), 18.43)

    # Одна пара пропущена целиком — это допустимо.
    start = tariffs.cell(2, _column(tariffs, "Действует с", 2), datetime(2022, 8, 10))
    start.number_format = "dd.mm.yyyy"
    tariffs.cell(2, _column(tariffs, "Тариф, руб./м² в месяц", 2), 20.59)

    tariffs.cell(2, _column(tariffs, "Действует с", 4), "08.2024")
    tariffs.cell(2, _column(tariffs, "Тариф, руб./м² в месяц", 4), "22,33")

    workbook.save(path)
    workbook.close()

    result = parse_house_data_template(path)

    assert [(row.start.input_label, row.amount) for row in result.tariffs] == [
        ("08.2020", "18,43"),
        ("10.08.2022", "20,59"),
        ("08.2024", "22,33"),
    ]

    repository.close()


def test_old_vertical_tariff_rows_still_import(tmp_path) -> None:
    repository = Repository(tmp_path / "app.db")
    organization = repository.create_organization(
        name="УК Пример",
        ogrn="1022600000102",
    )
    house = repository.create_house(
        organization_id=organization.id,
        address="г Ставрополь, ул Братская, д. 10",
    )

    path = tmp_path / "vertical-house-data.xlsx"
    write_house_data_template(
        path,
        houses=[house],
        tariffs_by_house={house.id: []},
    )

    workbook = load_workbook(path)
    del workbook["Тарифы"]
    tariffs = workbook.create_sheet("Тарифы")

    # Настоящий формат dev23, независимо от текущего порядка столбцов экспорта.
    tariffs.append(["_house_id", "Адрес", "Действует с", "Тариф, руб./м² в месяц"])
    tariffs.append([house.id, house.address, "08.2025", "22"])
    tariffs.append([house.id, house.address, "08.2026", "23"])

    workbook.save(path)
    workbook.close()

    result = parse_house_data_template(path)

    assert [(row.start.input_label, row.amount) for row in result.tariffs] == [
        ("08.2025", "22"),
        ("08.2026", "23"),
    ]

    repository.close()
