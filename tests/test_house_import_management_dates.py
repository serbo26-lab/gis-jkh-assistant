"""Management dates imported through real XLSX and isolated desktop workflows."""

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QDialog, QFileDialog, QMessageBox

import test_ui_workflow as workflows
from gis_jkh_assistant.app.house_import_preview import HouseImportPreview
from gis_jkh_assistant.gis_files.house_data import write_house_data_template
from gis_jkh_assistant.gis_files.house_extra_data import HOUSE_COLUMNS

desktop = workflows.desktop
qt_app = workflows.qt_app

STORED_DATES = {
    "management_agreement_date": "01.06.2018",
    "management_start_date": "02.07.2018",
    "management_end_date": "01.08.2018",
}
NEW_DATES = {
    "management_agreement_date": "2018-06-02",
    "management_start_date": "2018-07-03",
    "management_end_date": "2018-08-02",
}


def _house(desktop):
    desktop.repository.update_house(
        desktop.house.id, **STORED_DATES, management_data_source="gis",
    )
    return desktop.repository.get_house(desktop.house.id)


def _workbook(tmp_path, desktop, house, fields=None):
    path = tmp_path / "management-dates.xlsx"
    write_house_data_template(
        path, houses=[house], tariffs_by_house={}, repository=desktop.repository,
    )
    if fields is not None:
        book = load_workbook(path)
        sheet = book["Дома"]
        columns = {cell.value: cell.column for cell in sheet[1]}
        for cell in sheet[2][1:]:
            cell.value = None
        for field, value in fields.items():
            sheet.cell(2, columns[HOUSE_COLUMNS[field]], value)
        book.save(path)
        book.close()
    return path


def _dialogs(monkeypatch, path):
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *_args: (str(path), ""))
    errors, messages, previews = [], [], []
    for kind in ("critical", "warning"):
        monkeypatch.setattr(QMessageBox, kind, lambda *_args: errors.append(_args[2]))
    monkeypatch.setattr(QMessageBox, "information", lambda *_args: messages.append(_args[1]))

    def accept(dialog):
        previews.append([
            tuple(dialog.table.item(row, column).text() for column in range(4))
            for row in range(dialog.table.rowCount())
        ])
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(HouseImportPreview, "exec", accept)
    return errors, messages, previews


def test_same_dates_in_display_and_iso_formats_do_not_change_source_or_database(
    desktop, tmp_path, monkeypatch,
):
    house = _house(desktop)
    path = _workbook(tmp_path, desktop, house)
    errors, messages, previews = _dialogs(monkeypatch, path)
    before = list(desktop.repository.connection.iterdump())

    for _ in range(2):
        desktop.window.houses_page._import_house_data()
        assert messages[-1] == "Изменений нет"
        assert list(desktop.repository.connection.iterdump()) == before

    assert not errors
    assert not previews
    updated = desktop.repository.get_house(house.id)
    assert updated.management_data_source == "gis"
    assert {field: getattr(updated, field) for field in STORED_DATES} == STORED_DATES


@pytest.mark.parametrize("field", STORED_DATES)
def test_real_management_date_change_is_applied_once_and_marks_source_manual(
    desktop, tmp_path, monkeypatch, field,
):
    house = _house(desktop)
    path = _workbook(tmp_path, desktop, house, {field: NEW_DATES[field]})
    errors, messages, previews = _dialogs(monkeypatch, path)

    desktop.window.houses_page._import_house_data()

    assert not errors
    assert len(previews) == 1
    assert previews[0] == [(house.address, HOUSE_COLUMNS[field], STORED_DATES[field], NEW_DATES[field])]
    updated = desktop.repository.get_house(house.id)
    assert updated.management_data_source == "manual"
    assert {name: getattr(updated, name) for name in STORED_DATES} == {
        **STORED_DATES, field: NEW_DATES[field],
    }
    after_first = list(desktop.repository.connection.iterdump())

    desktop.window.houses_page._import_house_data()

    assert not errors
    assert len(previews) == 1
    assert messages[-1] == "Изменений нет"
    assert list(desktop.repository.connection.iterdump()) == after_first


def test_unrelated_house_change_preserves_equivalent_management_dates_and_source(
    desktop, tmp_path, monkeypatch,
):
    house = _house(desktop)
    path = _workbook(tmp_path, desktop, house, {**STORED_DATES, "comment": "Проверено"})
    errors, _, previews = _dialogs(monkeypatch, path)

    desktop.window.houses_page._import_house_data()

    assert not errors
    assert previews == [[(house.address, HOUSE_COLUMNS["comment"], "Не указано", "Проверено")]]
    updated = desktop.repository.get_house(house.id)
    assert updated.comment == "Проверено"
    assert updated.management_data_source == "gis"
    assert {field: getattr(updated, field) for field in STORED_DATES} == STORED_DATES


@pytest.mark.parametrize("fields", [
    {"management_end_date": "2018-06-01"},
    {"management_start_date": "2018-09-02"},
])
def test_management_end_before_start_is_rejected_with_mixed_stored_and_imported_formats(
    desktop, tmp_path, monkeypatch, fields,
):
    house = _house(desktop)
    path = _workbook(tmp_path, desktop, house, fields)
    errors, _, previews = _dialogs(monkeypatch, path)
    before = list(desktop.repository.connection.iterdump())

    desktop.window.houses_page._import_house_data()

    assert any("Начало управления позже окончания" in error for error in errors)
    assert not previews
    assert list(desktop.repository.connection.iterdump()) == before


@pytest.mark.parametrize("field", STORED_DATES)
def test_valid_import_replaces_invalid_legacy_management_date(
    desktop, tmp_path, monkeypatch, field,
):
    house = _house(desktop)
    path = _workbook(tmp_path, desktop, house, {field: NEW_DATES[field]})
    desktop.repository.update_house(house.id, **{field: "ошибочная дата"})
    errors, _, previews = _dialogs(monkeypatch, path)

    desktop.window.houses_page._import_house_data()

    assert not errors
    assert len(previews) == 1
    updated = desktop.repository.get_house(house.id)
    assert getattr(updated, field) == NEW_DATES[field]
    assert updated.management_data_source == "manual"


@pytest.mark.parametrize("field", STORED_DATES)
def test_invalid_imported_management_date_is_rejected_without_writes(
    desktop, tmp_path, monkeypatch, field,
):
    house = _house(desktop)
    path = _workbook(tmp_path, desktop, house, {field: "31.02.2018"})
    errors, _, previews = _dialogs(monkeypatch, path)
    before = list(desktop.repository.connection.iterdump())

    desktop.window.houses_page._import_house_data()

    assert any("Некорректная дата" in error for error in errors)
    assert not previews
    assert list(desktop.repository.connection.iterdump()) == before
