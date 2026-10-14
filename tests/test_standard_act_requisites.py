"""Historical signatories, workbook roundtrips and immutable generated acts."""
import json

import pytest
from openpyxl import load_workbook

from act_test_data import populate_act_requisites
from test_gis_downstream_sync import _setup
from gis_jkh_assistant.acceptance_act_details import address_parts, resolve_act_details, suggested_genitive
from gis_jkh_assistant.database.act_requisites_store import ActRequisitesStore
from gis_jkh_assistant.database.quality_assessment_store import QualityAssessmentStore
from gis_jkh_assistant.gis_files.house_data import write_house_data_template, parse_house_data_template
from gis_jkh_assistant.quality_automation import generate_acceptance_act_for_period, ActReadinessError
from gis_jkh_assistant.standard_acceptance_pdf import money_in_words


@pytest.fixture
def source(tmp_path):
    repo, wl = _setup(tmp_path)
    populate_act_requisites(repo, wl)
    yield repo, wl
    repo.close()


def test_signatory_is_chosen_on_act_date_and_overlaps_require_a_choice(source):
    repo, wl = source
    store = ActRequisitesStore(repo)
    house = repo.get_house(wl.house_id)
    old = store.list_people(house_id=house.id)[0]
    store.save_person(dict(old, valid_to="2026-01-31"), person_id=old["id"], house_id=house.id)
    next_id = store.save_person(dict(old, full_name="Сидорова Анна Ивановна", valid_from="2026-02-01", valid_to=""), house_id=house.id)
    assert resolve_act_details(repo, house.id, "31.01.2026")["customer"]["id"] == old["id"]
    assert resolve_act_details(repo, house.id, "28.02.2026")["customer"]["id"] == next_id
    store.save_person(dict(old, full_name="Смирнова Анна Ивановна", valid_from="2026-02-01", valid_to=""), house_id=house.id)
    assert "Выберите одного" in " ".join(resolve_act_details(repo, house.id, "28.02.2026")["issues"])
    assert not resolve_act_details(repo, house.id, "28.02.2026", customer_id=next_id)["issues"]


def test_generated_act_snapshot_and_bytes_survive_changes_to_live_requisites(source):
    repo, wl = source
    house = repo.get_house(wl.house_id)
    store = QualityAssessmentStore(repo)
    period = store.get_or_create_period(work_list_id=wl.id, year=2026, month=1)
    item = store.list_items(period.id)[0]
    store.update_item(item.id, included=True, cost=None, days_count=None, cost_reduction=None)
    fixation = repo.get_or_create_work_fixation_period(work_list_id=wl.id, year=2026, month=1)
    repo.update_work_fixation_item(repo.list_work_fixation_items(fixation.id)[0].id,
        fact_price="100", fact_amount="2", fact_count=1, fact_total_cost="200")
    generate_acceptance_act_for_period(store, period_id=period.id, organization_name="TEST", house_address="TEST")
    act = store.list_period_acts(period.id)[0]
    requisites = ActRequisitesStore(repo)
    snapshot = requisites.snapshot(act.id)
    assert snapshot["work_rows"][0]["unit_cost"] == "100"
    before = act.content
    director = requisites.list_people(organization_id=house.organization_id)[0]
    requisites.save_person(dict(director, full_name="Другой Подписант"), organization_id=house.organization_id, person_id=director["id"])
    repo.update_house(house.id, management_agreement_number="CHANGED")
    assert requisites.snapshot(act.id) == snapshot
    assert store.get_act(act.id).content == before
    assert snapshot["agreement_number"] == "TEST-1"
    assert snapshot["executor"]["full_name"] == "Иванов Иван Иванович"


def test_missing_requisites_never_replace_previous_act(source):
    repo, wl = source
    store = QualityAssessmentStore(repo)
    period = store.get_or_create_period(work_list_id=wl.id, year=2026, month=1)
    item = store.list_items(period.id)[0]
    store.update_item(item.id, included=True, cost=None, days_count=None, cost_reduction=None)
    previous = store.add_act(item_id=item.id, act_type="ACCEPTANCE_ACT", number="1", reg_date="31.01.2026",
        period_from="01.01.2026", period_to="31.01.2026", file_name="old.pdf", description=None,
        content_type="application/pdf", content=b"OLD DOCUMENT", generated=True)
    before = list(repo.connection.iterdump())
    with pytest.raises(ActReadinessError, match="фактическая стоимость"):
        generate_acceptance_act_for_period(store, period_id=period.id, organization_name="TEST", house_address="TEST")
    assert list(repo.connection.iterdump()) == before
    assert store.get_act(previous.id).content == b"OLD DOCUMENT"


def test_extended_workbook_roundtrip_keeps_house_data_and_person_history(source, tmp_path):
    repo, wl = source
    house = repo.get_house(wl.house_id)
    repo.update_house(house.id, residential_premises_count=17, non_residential_premises_count=0,
        residential_area="1234,5", non_residential_area="0", charge_area="1234,5",
        management_start_date="2021-04-03", management_end_date="2030-05-01")
    house = repo.get_house(house.id)
    file = tmp_path / "expanded.xlsx"
    write_house_data_template(file, houses=[house], tariffs_by_house={}, repository=repo)
    data = parse_house_data_template(file)
    fields = data.houses[0].fields
    assert fields["residential_premises_count"] == 17
    assert fields["non_residential_premises_count"] == 0
    assert fields["management_agreement_number"] == "TEST-1"
    assert fields["management_start_date"] == "2021-04-03"
    assert fields["management_end_date"] == "2030-05-01"
    assert len(data.people) == 2
    assert data.areas[0].non_residential_area == "0"
    assert data.people[0].fields["valid_from"] == "2020-01-01"
    assert data.people[1].ogrn == repo.get_organization(house.organization_id).ogrn
    workbook = load_workbook(file)
    sheet = workbook["Дома"]
    sheet.cell(2, 3, "-1")
    workbook.save(file)
    workbook.close()
    with pytest.raises(ValueError, match="Количество помещений"):
        parse_house_data_template(file)


def test_legacy_workbook_without_extra_sheets_is_supported(source, tmp_path):
    repo, wl = source
    file = tmp_path / "legacy.xlsx"
    write_house_data_template(file, houses=[repo.get_house(wl.house_id)], tariffs_by_house={})
    book = load_workbook(file)
    for sheet in ("Дома", "Представители домов", "Подписанты УО"):
        del book[sheet]
    book.save(file)
    book.close()
    parsed = parse_house_data_template(file)
    assert not parsed.houses and not parsed.people


@pytest.mark.parametrize("value,expected", [
    ("0", "Ноль рублей 00 копеек"), ("1.01", "Один рубль 01 копейка"),
    ("2002.12", "Две тысячи два рубля 12 копеек"),
    ("1000000", "Один миллион рублей 00 копеек"),
])
def test_money_in_words(value, expected):
    assert money_in_words(value) == expected


def test_compound_house_address_is_not_duplicated():
    parts = address_parts("355000, г. Тестовый, ул. Учебная, д. 15А, корп. 2, стр. 3")
    assert parts == dict(locality="г. Тестовый", house_number="15А, корп. 2, стр. 3",
                         address_without_number="355000, г. Тестовый, ул. Учебная")
    assert address_parts("Неоднозначный адрес 15")["house_number"] == ""


@pytest.mark.parametrize("address,number,remaining", [
    ("г Москва, ул Учебная д. 15А", "15А", "г Москва, ул Учебная"),
    ("с. Тестовое, ул. Учебная, дом 7/2, строение 4", "7/2, строение 4", "с. Тестовое, ул. Учебная"),
    ("г. Тестовый, ул. Учебная, д. 7, неизвестное уточнение", "", ""),
    ("г. Тестовый, ул. Учебная 7", "", ""),
])
def test_explicit_address_markers_preserve_compound_and_ambiguous_addresses(address, number, remaining):
    parts = address_parts(address)
    assert parts["house_number"] == number
    assert parts["address_without_number"] == remaining


def test_long_address_suffixes_do_not_require_regex_backtracking():
    suffixes = ", корпус 2" * 2000
    parts = address_parts("г. Тестовый, ул. Учебная, д. 15А" + suffixes)
    assert parts["house_number"] == "15А" + suffixes
    assert parts["address_without_number"] == "г. Тестовый, ул. Учебная"
    assert address_parts("д." + " " * 100000)["house_number"] == ""


def test_fio_suggestion_keeps_indeclinable_family_name():
    assert suggested_genitive("Шевченко Александр Андреевич", "male") == "Шевченко Александра Андреевича"
    assert suggested_genitive("Иванова Мария Ивановна", "female") == "Ивановой Марии Ивановны"


def test_snapshot_requires_an_existing_act(source):
    repo, wl = source
    QualityAssessmentStore(repo)
    import sqlite3
    with pytest.raises(sqlite3.IntegrityError):
        ActRequisitesStore(repo).save_snapshot("no-such-act", json.loads('{"test": true}'))


def test_legacy_invalid_house_dates_are_actionable_readiness_issues(source):
    repo, wl = source
    # Simulate a legacy record, not a new validated form submission.
    repo.connection.execute("UPDATE houses SET management_agreement_date=? WHERE id=?", ("ошибка даты", wl.house_id))
    details = resolve_act_details(repo, wl.house_id, "31.01.2026")
    assert details["agreement_date"] == ""
    assert any("Дата договора управления" in issue for issue in details["issues"])


def test_snapshot_failure_rolls_back_replacement_and_preserves_signatory_images(source, monkeypatch):
    repo, wl = source
    house = repo.get_house(wl.house_id)
    people = ActRequisitesStore(repo)
    director = people.list_people(organization_id=house.organization_id)[0]
    people.save_signature(director["id"], b"TEST SIGNATURE", "test.png")
    people.save_options(house.organization_id, stamp=b"TEST STAMP", stamp_name="stamp.png", insert_images=0)
    store = QualityAssessmentStore(repo)
    period = store.get_or_create_period(work_list_id=wl.id, year=2026, month=1)
    item = store.list_items(period.id)[0]
    store.update_item(item.id, included=True, cost=None, days_count=None, cost_reduction=None)
    fixation = repo.get_or_create_work_fixation_period(work_list_id=wl.id, year=2026, month=1)
    repo.update_work_fixation_item(repo.list_work_fixation_items(fixation.id)[0].id,
        fact_price="100", fact_amount="2", fact_count=1, fact_total_cost="200")
    generate_acceptance_act_for_period(store, period_id=period.id, organization_name="TEST", house_address="TEST")
    act = store.list_period_acts(period.id)[0]
    before = list(repo.connection.iterdump())
    def fail(*_):
        raise RuntimeError("TEST SNAPSHOT FAILURE")
    monkeypatch.setattr(ActRequisitesStore, "save_snapshot", fail)
    with pytest.raises(RuntimeError, match="TEST SNAPSHOT FAILURE"):
        generate_acceptance_act_for_period(store, period_id=period.id, organization_name="TEST", house_address="TEST")
    assert list(repo.connection.iterdump()) == before
    snapshot = people.snapshot(act.id)
    people.save_signature(director["id"], b"NEW SIGNATURE", "new.png")
    people.save_options(house.organization_id, stamp=b"NEW STAMP")
    assert people.snapshot(act.id) == snapshot
    assert store.get_act(act.id).content == act.content
