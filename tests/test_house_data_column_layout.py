"""Copying a rectangular tariff history must not overwrite hidden house IDs."""

from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from gis_jkh_assistant.database import Repository
from gis_jkh_assistant.database.act_requisites_store import ActRequisitesStore
from gis_jkh_assistant.gis_files.house_data import (
    HouseDataImportError,
    parse_house_data_template,
    write_house_data_template,
)


@pytest.fixture
def houses(tmp_path):
    repository = Repository(tmp_path / "layout.db")
    organization = repository.create_organization(name="Вымышленная УО", ogrn="1022600000001")
    records = [
        repository.create_house(
            organization_id=organization.id,
            address=f"г. Тестовый, ул. Учебная, д. {number}",
            house_object_guid=f"fictional-object-{number}",
            fias_house_guid=f"fictional-fias-{number}",
            gis_unique_house_number=f"fictional-unique-{number}",
        )
        for number in (1, 2, 3)
    ]
    yield repository, records
    repository.close()


def _column(sheet, header):
    return next(cell.column for cell in sheet[1] if cell.value == header)


def test_all_house_sheets_keep_technical_columns_before_the_visible_block(houses, tmp_path):
    repository, records = houses
    path = tmp_path / "template.xlsx"
    write_house_data_template(path, houses=records, tariffs_by_house={}, repository=repository)
    workbook = load_workbook(path)
    try:
        expected = {
            "Площади": ("Адрес", 5, "F2"),
            "Тарифы": ("Адрес", 5, "F2"),
            "Дома": ("Адрес", 5, "F2"),
            "Представители домов": ("Адрес", 6, "G2"),
            "Подписанты УО": ("УО", 3, "D2"),
        }
        for name, (header, first_visible, freeze) in expected.items():
            sheet = workbook[name]
            assert _column(sheet, header) == first_visible
            assert sheet.freeze_panes == freeze
            for cell in sheet[1]:
                hidden = bool(sheet.column_dimensions[cell.column_letter].hidden)
                assert hidden == (cell.column < first_visible), (name, cell.coordinate)
                assert str(cell.value).startswith("_") == hidden
        tariffs = workbook["Тарифы"]
        assert sum(cell.value == "Действует с" for cell in tariffs[1]) >= 20
    finally:
        workbook.close()


@pytest.mark.parametrize("pair_count", [16, 23])
def test_rectangular_multi_house_paste_keeps_every_tariff_and_identity(houses, tmp_path, pair_count):
    repository, records = houses
    path = tmp_path / "pasted.xlsx"
    write_house_data_template(path, houses=records, tariffs_by_house={}, repository=repository)
    workbook = load_workbook(path)
    sheet = workbook["Тарифы"]
    address_column = _column(sheet, "Адрес")
    identities = {
        row: tuple(sheet.cell(row, column).value for column in range(1, address_column))
        for row in range(2, len(records) + 2)
    }
    expected = []
    for row, house in enumerate(records, 2):
        # A normal clipboard block contains the address followed by uninterrupted
        # date/amount pairs, just like a yearly tariff spreadsheet.
        clipboard = [house.address]
        for index in range(pair_count):
            year = 2010 + index
            day = 1 if index % 2 == 0 else 10
            start = datetime(year, 8, day)
            amount = Decimal("18.40") + Decimal(index) / 100 + Decimal(row - 2)
            clipboard.extend((start, str(amount).replace(".", ",")))
            expected.append((
                house.id,
                house.address,
                start.date().isoformat(),
                "month" if day == 1 else "day",
                format(amount, "f").rstrip("0").rstrip(".").replace(".", ","),
                house.house_object_guid,
                house.fias_house_guid,
                house.gis_unique_house_number,
            ))
        for offset, value in enumerate(clipboard):
            sheet.cell(row, address_column + offset, value)
        for index in range(pair_count):
            sheet.cell(row, address_column + 1 + index * 2).number_format = (
                "mm.yyyy" if index % 2 == 0 else "dd.mm.yyyy"
            )
    workbook.save(path)
    workbook.close()

    imported = parse_house_data_template(path)
    assert [(
        row.house_id, row.address, row.start.value.isoformat(), row.start.precision,
        row.amount, row.house_object_guid, row.fias_house_guid, row.gis_unique_house_number,
    ) for row in imported.tariffs] == expected
    workbook = load_workbook(path)
    try:
        sheet = workbook["Тарифы"]
        assert {
            row: tuple(sheet.cell(row, column).value for column in range(1, address_column))
            for row in identities
        } == identities
    finally:
        workbook.close()


@pytest.mark.parametrize("include_identities", [False, True])
def test_old_horizontal_tariffs_remain_importable(houses, tmp_path, include_identities):
    _, records = houses
    house = records[0]
    workbook = Workbook()
    areas = workbook.active
    areas.title = "Площади"
    areas.append(["_house_id", "Адрес", "Площадь к начислению, м²"])
    areas.append([house.id, house.address, "0"])
    tariffs = workbook.create_sheet("Тарифы")
    header = ["_house_id", "Адрес", "Действует с", "Тариф, руб./м² в месяц",
              "Действует с", "Тариф, руб./м² в месяц"]
    row = [house.id, house.address, "08.2025", "22,15", "10.08.2026", "23"]
    if include_identities:
        header.extend(["_house_object_guid", "_fias_house_guid", "_gis_unique_house_number"])
        row.extend([house.house_object_guid, house.fias_house_guid, house.gis_unique_house_number])
    tariffs.append(header)
    tariffs.append(row)
    path = tmp_path / "old-horizontal.xlsx"
    workbook.save(path)
    workbook.close()
    imported = parse_house_data_template(path)
    assert imported.areas[0].area == "0"
    assert [(row.start.input_label, row.amount) for row in imported.tariffs] == [
        ("08.2025", "22,15"), ("10.08.2026", "23"),
    ]
    if include_identities:
        assert imported.tariffs[0].house_object_guid == house.house_object_guid


def test_old_extra_sheet_order_keeps_house_fields_zeros_and_person_ids(houses, tmp_path):
    repository, records = houses
    house = records[0]
    repository.update_house(house.id, residential_premises_count=17, non_residential_premises_count=0,
                            management_agreement_number="FICTIONAL-1", management_start_date="2021-04-03")
    store = ActRequisitesStore(repository)
    store.save_person(dict(full_name="Иванов Иван Иванович", position="Директор", role="Руководитель",
                           valid_from="2020-01-01", basis_kind="Устав"), organization_id=house.organization_id)
    store.save_person(dict(full_name="Петрова Мария Ивановна", role="Председатель совета МКД", premises="7",
                           valid_from="2020-01-01", basis_kind="Решение ОСС", basis_number="1",
                           basis_date="2020-01-01"), house_id=house.id)
    path = tmp_path / "legacy-extra.xlsx"
    write_house_data_template(path, houses=[repository.get_house(house.id)], tariffs_by_house={}, repository=repository)
    before = parse_house_data_template(path)
    workbook = load_workbook(path)
    legacy_prefix = {
        "Дома": ["_house_id", "Адрес"],
        "Представители домов": ["_house_id", "Адрес", "_person_id"],
        "Подписанты УО": ["_organization_id", "УО", "ОГРН УО", "_person_id"],
    }
    for name, prefix in legacy_prefix.items():
        sheet = workbook[name]
        values = list(sheet.values)
        headers = list(values[0])
        visible = [header for header in headers if not header.startswith("_") and header not in prefix]
        trailing = [header for header in headers if header.startswith("_") and header not in prefix]
        order = prefix + visible + trailing
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(order)
        for row in values[1:]:
            sheet.append([row[headers.index(header)] for header in order])
    workbook.save(path)
    workbook.close()
    after = parse_house_data_template(path)
    assert after.houses == before.houses
    assert after.people == before.people
    assert after.houses[0].fields["non_residential_premises_count"] == 0
    assert all(person.person_id for person in after.people)


@pytest.mark.parametrize("malformation", [
    "renamed_header", "trailing_unknown_header", "incomplete_trailing_pair", "missing_named_amount", "invalid_named_amount",
])
def test_invalid_tariff_columns_are_reported_instead_of_silently_ignored(houses, tmp_path, malformation):
    repository, records = houses
    path = tmp_path / "invalid.xlsx"
    write_house_data_template(path, houses=records, tariffs_by_house={}, repository=repository)
    workbook = load_workbook(path)
    sheet = workbook["Тарифы"]
    date_column = _column(sheet, "Действует с")
    sheet.cell(2, date_column, "08.2020")
    sheet.cell(2, date_column + 1, "18,43")
    if malformation == "renamed_header":
        sheet.cell(1, date_column + 3, "Неизвестный столбец")
    elif malformation == "missing_named_amount":
        sheet.cell(2, date_column + 1).value = None
    elif malformation == "invalid_named_amount":
        sheet.cell(2, date_column + 1, "некорректный тариф")
    else:
        trailing = sheet.max_column + 1
        if malformation == "trailing_unknown_header":
            sheet.cell(1, trailing, "Неизвестный столбец")
            sheet.cell(2, trailing, "08.2040")
            sheet.cell(2, trailing + 1, "25,01")
        else:
            sheet.cell(2, trailing, "08.2040")
    workbook.save(path)
    workbook.close()
    with pytest.raises(HouseDataImportError):
        parse_house_data_template(path)
