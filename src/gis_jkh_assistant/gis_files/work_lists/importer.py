"""Import GIS ЖКХ work-list XLSX files as reusable user templates."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class WorkTemplateImportError(ValueError):
    """Raised when a workbook is not recognizable as a work-list template."""


@dataclass(frozen=True, slots=True)
class ImportedWorkItem:
    position: int
    work_code: str | None
    work_name: str
    work_guid: str | None
    price: str | None = None
    volume: str | None = None
    quantity: str | None = None
    cost: str | None = None


@dataclass(frozen=True, slots=True)
class ImportedWorkTemplate:
    file_type: str | None
    version: str | None
    house_guid: str | None
    start_year: int | None
    start_month: int | None
    end_year: int | None
    end_month: int | None
    items: tuple[ImportedWorkItem, ...]


def parse_work_template(path: str | Path) -> ImportedWorkTemplate:
    """Read either a blank GIS template or a previously filled work-list XLSX."""

    workbook = load_workbook(Path(path), read_only=False, data_only=False)
    try:
        if "Перечень работ и услуг" not in workbook.sheetnames:
            raise WorkTemplateImportError(
                "В книге не найден лист «Перечень работ и услуг»."
            )
        if "СпрРабУсл" not in workbook.sheetnames:
            raise WorkTemplateImportError(
                "В книге не найден справочник работ «СпрРабУсл»."
            )

        lookup = _read_work_lookup(workbook["СпрРабУсл"])
        items = _read_visible_items(workbook["Перечень работ и услуг"], lookup)
        if not items:
            items = [
                ImportedWorkItem(
                    position=index,
                    work_code=code,
                    work_name=name,
                    work_guid=guid,
                )
                for index, (label, guid) in enumerate(lookup.items(), start=1)
                for code, name in [_split_work_label(label)]
            ]

        file_type, version = _read_conf(workbook)
        house_guid, start_year, start_month, end_year, end_month = _read_options(workbook)

        return ImportedWorkTemplate(
            file_type=file_type,
            version=version,
            house_guid=house_guid,
            start_year=start_year,
            start_month=start_month,
            end_year=end_year,
            end_month=end_month,
            items=tuple(items),
        )
    finally:
        workbook.close()


def _read_work_lookup(sheet: Any) -> dict[str, str | None]:
    result: dict[str, str | None] = {}
    for row in sheet.iter_rows(values_only=True):
        label = _text(row[0] if row else None)
        if not label:
            continue
        guid = _text(row[1] if len(row) > 1 else None)
        result[label] = guid
    if not result:
        raise WorkTemplateImportError("Справочник работ в файле пуст.")
    return result


def _read_visible_items(sheet: Any, lookup: dict[str, str | None]) -> list[ImportedWorkItem]:
    result: list[ImportedWorkItem] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = list(row) + [None] * 6
        label = _text(values[0])
        if not label:
            continue
        code, name = _split_work_label(label)
        result.append(
            ImportedWorkItem(
                position=len(result) + 1,
                work_code=code,
                work_name=name,
                work_guid=lookup.get(label),
                price=_text(values[1]),
                volume=_text(values[2]),
                quantity=_text(values[3]),
                cost=_text(values[4]),
            )
        )
    return result


def _read_conf(workbook: Any) -> tuple[str | None, str | None]:
    if "conf" not in workbook.sheetnames:
        return None, None
    values: dict[str, str] = {}
    for row in workbook["conf"].iter_rows(values_only=True):
        if len(row) < 2:
            continue
        key = _text(row[0])
        value = _text(row[1])
        if key and value:
            values[key.upper()] = value
    return values.get("TYPE"), values.get("VERSION")


def _read_options(
    workbook: Any,
) -> tuple[str | None, int | None, int | None, int | None, int | None]:
    if "ОпцииПеречня" not in workbook.sheetnames:
        return None, None, None, None, None

    rows = list(workbook["ОпцииПеречня"].iter_rows(values_only=True))
    house_guid = _text(rows[0][1]) if len(rows) > 0 and len(rows[0]) > 1 else None
    start_year = _int_value(rows[1][1]) if len(rows) > 1 and len(rows[1]) > 1 else None
    start_month = _int_value(rows[2][1]) if len(rows) > 2 and len(rows[2]) > 1 else None
    end_year = _int_value(rows[3][1]) if len(rows) > 3 and len(rows[3]) > 1 else None
    end_month = _int_value(rows[4][1]) if len(rows) > 4 and len(rows[4]) > 1 else None
    return house_guid, start_year, start_month, end_year, end_month


def _split_work_label(label: str) -> tuple[str | None, str]:
    before, dot, after = label.partition(".")
    if dot and before.strip().isdigit() and after.strip():
        return before.strip(), after.strip()
    return None, label.strip()


def _int_value(value: Any) -> int | None:
    if value is None:
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None
