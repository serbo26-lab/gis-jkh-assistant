"""Optional workbook sheets for house requisites and dated act signatories."""
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from gis_jkh_assistant.acceptance_act_details import address_parts, display_date, iso_date
from gis_jkh_assistant.database.act_requisites_store import ActRequisitesStore, normalized_person


HOUSE_COLUMNS = {
    "residential_premises_count": "Количество жилых помещений",
    "non_residential_premises_count": "Количество нежилых помещений",
    "management_agreement_number": "Номер договора управления",
    "management_agreement_date": "Дата договора управления",
    "management_start_date": "Начало управления",
    "management_end_date": "Окончание управления (первый день вне управления)",
    "cadastral_number": "Кадастровый номер", "comment": "Комментарий",
    "locality": "Населённый пункт для акта", "house_number": "Номер дома для акта",
    "address_without_number": "Адрес для акта без номера дома",
}
PERSON_COLUMNS = {
    "full_name": "ФИО", "name_genitive": "ФИО в родительном падеже",
    "position": "Должность", "position_genitive": "Должность в родительном падеже",
    "role": "Статус представителя", "premises": "Квартира / помещение",
    "gender": "Согласование (male / female)", "basis_kind": "Основание полномочий",
    "basis_number": "Номер основания", "basis_date": "Дата основания", "basis_text": "Основание полным текстом",
    "valid_from": "Начало полномочий", "valid_to": "Окончание полномочий (включительно)",
}
IDENTITIES = ("_house_object_guid", "_fias_house_guid", "_gis_unique_house_number")


@dataclass
class ExtraHouseRow:
    house_id: str | None
    address: str
    row_number: int
    fields: dict = field(default_factory=dict)
    house_object_guid: str | None = None
    fias_house_guid: str | None = None
    gis_unique_house_number: str | None = None


@dataclass
class PersonRow(ExtraHouseRow):
    person_id: str | None = None
    ogrn: str = ""


def _style(sheet):
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        sheet.column_dimensions[cell.column_letter].width = 29
        if str(cell.value).startswith("_"):
            sheet.column_dimensions[cell.column_letter].hidden = True
    sheet.column_dimensions["B"].width = 58
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "@"


def write_extra_sheets(workbook, houses, repository=None):
    store = ActRequisitesStore(repository) if repository else None
    house_sheet = workbook.create_sheet("Дома")
    house_sheet.append(["_house_id", "Адрес", *HOUSE_COLUMNS.values(), *IDENTITIES])
    people_sheet = workbook.create_sheet("Представители домов")
    people_sheet.append(["_house_id", "Адрес", "_person_id", *PERSON_COLUMNS.values(), *IDENTITIES])
    org_sheet = workbook.create_sheet("Подписанты УО")
    org_sheet.append(["_organization_id", "УО", "ОГРН УО", "_person_id", *PERSON_COLUMNS.values()])
    for house in houses:
        parts = address_parts(house.address)
        if store:
            parts.update({k: v for k, v in store.house_fields(house.id).items() if v})
        identity = [getattr(house, key[1:], None) for key in IDENTITIES]
        values = []
        for key in HOUSE_COLUMNS:
            value = parts.get(key, getattr(house, key, None))
            values.append(display_date(value) if key.endswith("_date") and value else value)
        house_sheet.append([house.id, house.address, *values, *identity])
        representatives = store.list_people(house_id=house.id) if store else []
        for person in representatives or [{}]:
            people_sheet.append([house.id, house.address, person.get("id"),
                *[display_date(person[k]) if k in {"basis_date", "valid_from", "valid_to"} and person.get(k)
                  else person.get(k) for k in PERSON_COLUMNS], *identity])
    if repository:
        for organization_id in dict.fromkeys(house.organization_id for house in houses):
            organization = repository.get_organization(organization_id)
            for person in store.list_people(organization_id=organization_id) or [{}]:
                org_sheet.append([organization_id, organization.name, organization.ogrn, person.get("id"),
                    *[display_date(person[k]) if k in {"basis_date", "valid_from", "valid_to"} and person.get(k)
                      else person.get(k) for k in PERSON_COLUMNS]])
    for sheet in (house_sheet, people_sheet, org_sheet):
        _style(sheet)
    for sheet in (people_sheet, org_sheet):
        headers = [c.value for c in sheet[1]]
        roles = '"Руководитель,Уполномоченный сотрудник"' if sheet == org_sheet else '"Председатель совета МКД,Уполномоченный собственник"'
        bases = '"Устав,Доверенность,Иное"' if sheet == org_sheet else '"Решение ОСС,Доверенность,Иное"'
        for header, formula in (("Статус представителя", roles), ("Основание полномочий", bases),
                                ("Согласование (male / female)", '"male,female"')):
            validation = DataValidation(type="list", formula1=formula, allow_blank=True)
            sheet.add_data_validation(validation)
            letter = sheet.cell(1, headers.index(header) + 1).column_letter
            validation.add(f"{letter}2:{letter}10000")


def parse_extra_sheets(workbook):
    houses, people = [], []
    for name in ("Дома", "Представители домов", "Подписанты УО"):
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        headers = [str(c.value or "").strip() for c in sheet[1]]
        required = {"_house_id", "Адрес"} if name != "Подписанты УО" else {"ОГРН УО"}
        if not required.issubset(headers):
            raise ValueError(f"Лист «{name}»: отсутствуют служебные столбцы или адрес / ОГРН.")
        for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            cells = dict(zip(headers, values))
            columns = HOUSE_COLUMNS if name == "Дома" else PERSON_COLUMNS
            fields = {key: cells[label] for key, label in columns.items()
                      if label in cells and cells[label] is not None and str(cells[label]).strip() != ""}
            if not fields:
                continue
            try:
                for key, value in list(fields.items()):
                    if key in {"management_agreement_date", "management_start_date", "management_end_date", "basis_date", "valid_from", "valid_to"}:
                        fields[key] = iso_date(value.date().isoformat() if isinstance(value, datetime)
                                               else value.isoformat() if isinstance(value, date) else value)
                    elif key.endswith("_premises_count"):
                        count = Decimal(str(value).replace(",", "."))
                        if not count.is_finite() or count < 0 or count != count.to_integral_value():
                            raise ValueError("Количество помещений должно быть целым неотрицательным числом.")
                        fields[key] = int(count)
                    else:
                        fields[key] = str(value).strip()
                identity = {key[1:]: str(cells.get(key) or "").strip() or None for key in IDENTITIES}
                basic = dict(house_id=str(cells.get("_house_id") or "").strip() or None,
                             address=str(cells.get("Адрес") or "").strip(), row_number=number,
                             fields=fields, **identity)
                if name == "Дома":
                    houses.append(ExtraHouseRow(**basic))
                else:
                    if name == "Подписанты УО" and not str(cells.get("ОГРН УО") or "").strip():
                        raise ValueError("Укажите ОГРН управляющей организации.")
                    people.append(PersonRow(**basic, person_id=str(cells.get("_person_id") or "").strip() or None,
                                            ogrn=str(cells.get("ОГРН УО") or "").strip()))
            except (ValueError, InvalidOperation) as exc:
                raise ValueError(f"Лист «{name}», строка {number}: {exc}") from exc
    return tuple(houses), tuple(people)


def prepare_person_change(repository, row, *, organization_id=None, house_id=None):
    store = ActRequisitesStore(repository)
    existing = store.list_people(organization_id=organization_id, house_id=house_id)
    matched = [p for p in existing if p["id"] == row.person_id] if row.person_id else []
    if row.person_id and not matched:
        # A workbook can survive database recreation. Reconcile by explicit
        # identity and dates, never by row position or a foreign local UUID.
        foreign = repository.connection.execute("SELECT id FROM act_people WHERE id=?", (row.person_id,)).fetchone()
        if foreign:
            raise ValueError("Служебный идентификатор подписанта принадлежит другому дому или УО.")
    if not matched:
        matched = [p for p in existing if p["full_name"].casefold() == str(row.fields.get("full_name", "")).casefold()
                   and p["valid_from"] == row.fields.get("valid_from")]
    if len(matched) > 1:
        raise ValueError("Несколько записей подписанта с одинаковыми ФИО и началом полномочий.")
    previous = matched[0] if matched else {}
    merged = {**previous, **row.fields}
    if "position" in row.fields and row.fields["position"] != previous.get("position"):
        if not row.fields.get("position_genitive") or row.fields["position_genitive"] == previous.get("position_genitive", ""):
            merged["position_genitive"] = ""
    data = normalized_person(merged)
    return previous.get("id"), data, previous
