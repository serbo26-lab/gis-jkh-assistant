"""Portable XLSX template for house charge areas and tariff history."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from gis_jkh_assistant.database.models import House, HouseTariff
from gis_jkh_assistant.period_dates import ParsedPeriodDate, parse_period_date


class HouseDataImportError(ValueError):
    """Raised when a house-data workbook is structurally invalid."""


@dataclass(frozen=True, slots=True)
class ImportedAreaRow:
    house_id: str | None
    address: str
    area: str
    row_number: int
    residential_area: str | None = None
    non_residential_area: str | None = None
    house_object_guid: str | None = None
    fias_house_guid: str | None = None
    gis_unique_house_number: str | None = None
    charge_explicit: bool = True


@dataclass(frozen=True, slots=True)
class ImportedTariffRow:
    house_id: str | None
    address: str
    start: ParsedPeriodDate
    amount: str
    row_number: int
    house_object_guid: str | None = None
    fias_house_guid: str | None = None
    gis_unique_house_number: str | None = None


@dataclass(frozen=True, slots=True)
class HouseDataImportResult:
    areas: tuple[ImportedAreaRow, ...]
    tariffs: tuple[ImportedTariffRow, ...]
    warnings: tuple[str, ...] = ()
    houses: tuple = ()
    people: tuple = ()


AREA_SHEET = "Площади"
TARIFF_SHEET = "Тарифы"
INSTRUCTION_SHEET = "Инструкция"
MIN_TARIFF_PAIRS = 10
IDENTITY_HEADERS = ["_house_object_guid", "_fias_house_guid", "_gis_unique_house_number"]
AREA_HEADERS = ["Площадь жилых помещений, м²", "Площадь нежилых помещений, м²"]


def write_house_data_template(
    path: str | Path,
    *,
    houses: Iterable[House],
    tariffs_by_house: dict[str, Iterable[HouseTariff]],
    repository=None,
) -> None:
    """Write a reusable workbook prefilled with the current houses."""

    workbook = Workbook()
    instruction = workbook.active
    instruction.title = INSTRUCTION_SHEET
    instruction.append(["Импорт площадей и тарифов"])
    instruction["A1"].font = Font(bold=True, size=14)
    instructions = [
        "Лист «Площади»: заполните жилую, нежилую и/или начисляемую площадь. Если начисляемая не указана, используется сумма указанных жилой и нежилой; ноль сохраняется как значение.",
        "Лист «Тарифы»: один дом занимает одну строку, история идет вправо парами "
        "«Действует с» / «Тариф». Пустые пары можно пропускать.",
        "Дата «08.2025» означает начало с 01.08.2025.",
        "Дата «10.08.2025» сохраняется как точная дата 10.08.2025.",
        "Конечную дату тарифа вводить не нужно: она определяется следующим изменением.",
        "Если тариф не менялся, повторять его каждый год не обязательно.",
        "Старый вертикальный формат (одна строка = одна точка тарифа) тоже поддерживается.",
        "Служебные идентификаторы скрыты. Адрес используется для проверки соответствия строк; после пересоздания базы проверьте сопоставление перед применением.",
    ]
    for line in instructions:
        instruction.append([line])
    instruction.column_dimensions["A"].width = 105

    areas = workbook.create_sheet(AREA_SHEET)
    areas.append(["_house_id", "Адрес", "Площадь к начислению, м²", *AREA_HEADERS, *IDENTITY_HEADERS])
    _style_header(areas)

    sorted_houses = sorted(houses, key=lambda item: item.address.casefold())
    histories = {
        house.id: sorted(
            list(tariffs_by_house.get(house.id, ())),
            key=lambda item: item.effective_from_date,
        )
        for house in sorted_houses
    }
    pair_count = max(
        MIN_TARIFF_PAIRS,
        max((len(items) for items in histories.values()), default=0),
    )

    tariffs = workbook.create_sheet(TARIFF_SHEET)
    tariff_header = ["_house_id", "Адрес"]
    for _ in range(pair_count):
        tariff_header.extend(["Действует с", "Тариф, руб./м² в месяц"])
    tariffs.append([*tariff_header, *IDENTITY_HEADERS])
    _style_header(tariffs)

    for house in sorted_houses:
        identity = [house.house_object_guid, house.fias_house_guid, house.gis_unique_house_number]
        areas.append([house.id, house.address, house.charge_area or "", house.residential_area or "", house.non_residential_area or "", *identity])

        tariff_row: list[object] = [house.id, house.address]
        existing = histories[house.id]
        for index in range(pair_count):
            if index < len(existing):
                tariff = existing[index]
                tariff_row.extend([tariff.effective_from_label, tariff.amount])
            else:
                tariff_row.extend(["", ""])
        tariffs.append([*tariff_row, *identity])

    for row in range(2, areas.max_row + 1):
        for column in (3, 4, 5):
            areas.cell(row, column).number_format = "@"
    for row in range(2, tariffs.max_row + 1):
        for column in range(3, tariffs.max_column + 1):
            tariffs.cell(row, column).number_format = "@"

    for sheet in (areas, tariffs):
        sheet.column_dimensions["A"].hidden = True
        for cell in sheet[1]:
            if cell.value in IDENTITY_HEADERS:
                sheet.column_dimensions[cell.column_letter].hidden = True
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions
    areas.column_dimensions["B"].width = 70
    for column in ("C", "D", "E"):
        areas.column_dimensions[column].width = 28
    tariffs.column_dimensions["B"].width = 70
    for column in range(3, tariffs.max_column + 1):
        tariffs.column_dimensions[
            tariffs.cell(1, column).column_letter
        ].width = 20 if column % 2 else 28

    from .house_extra_data import write_extra_sheets
    write_extra_sheets(workbook, sorted_houses, repository)
    instruction.append(["Лист «Дома»: помещения, договор, сроки управления и адресные реквизиты акта. Пустые ячейки не удаляют сохранённые значения; явный ноль меняет количество помещений на ноль."])
    instruction.append(["Листы «Представители домов» и «Подписанты УО»: одна строка — один период полномочий. Заполните ФИО, основание и точную дату начала; окончание необязательно и включается в период. Для новой записи скопируйте строку дома / УО и очистите _person_id. Подпись и печать загружаются только в программе."])
    workbook.save(Path(path))


def parse_house_data_template(path: str | Path) -> HouseDataImportResult:
    """Parse the application house-data workbook."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        missing = [name for name in (AREA_SHEET, TARIFF_SHEET) if name not in workbook.sheetnames]
        if missing:
            raise HouseDataImportError(
                "В файле отсутствуют обязательные листы: " + ", ".join(missing)
            )

        areas_sheet = workbook[AREA_SHEET]
        tariffs_sheet = workbook[TARIFF_SHEET]
        _require_headers(
            areas_sheet,
            ["_house_id", "Адрес", "Площадь к начислению, м²"],
        )
        tariff_pairs = _tariff_column_pairs(tariffs_sheet)

        area_headers = [str(c.value or "").strip() for c in areas_sheet[1]]
        tariff_headers = [str(c.value or "").strip() for c in tariffs_sheet[1]]
        def identity_fields(values, headers):
            return {header[1:]: _text(_at(values, headers.index(header)))
                    for header in IDENTITY_HEADERS if header in headers}
        area_rows: list[ImportedAreaRow] = []
        tariff_rows: list[ImportedTariffRow] = []
        warnings: list[str] = []

        for row_number, values in enumerate(
            areas_sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            house_id = _text(_at(values, 0))
            address = _text(_at(values, 1)) or ""
            raw_area = _at(values, 2)
            raw_components = [_at(values, area_headers.index(header)) if header in area_headers else None
                              for header in AREA_HEADERS]
            supplied = [v is not None and str(v).strip() != "" for v in (raw_area, *raw_components)]
            if not any(supplied):
                continue
            if not house_id and not address:
                raise HouseDataImportError(
                    f"Лист «{AREA_SHEET}», строка {row_number}: не указан дом."
                )
            try:
                components = [normalize_area(v) if present else None for v, present in zip(raw_components, supplied[1:])]
                area = normalize_area(raw_area) if supplied[0] else normalize_area(sum(
                    (Decimal(v.replace(",", ".")) for v in components if v is not None), Decimal(0)))
            except ValueError as exc:
                raise HouseDataImportError(
                    f"Лист «{AREA_SHEET}», строка {row_number}: {exc}"
                ) from exc
            area_rows.append(ImportedAreaRow(house_id, address, area, row_number, *components,
                                           **identity_fields(values, area_headers), charge_explicit=supplied[0]))

        seen_tariffs: set[tuple[str, str]] = set()
        for row_number, cells in enumerate(
            tariffs_sheet.iter_rows(min_row=2),
            start=2,
        ):
            house_id = _text(_cell_value(cells, 1))
            address = _text(_cell_value(cells, 2)) or ""

            for pair_number, (start_column, amount_column) in enumerate(
                tariff_pairs,
                start=1,
            ):
                start_cell = _cell(cells, start_column)
                amount_cell = _cell(cells, amount_column)
                raw_start = start_cell.value if start_cell is not None else None
                raw_amount = amount_cell.value if amount_cell is not None else None
                start_blank = raw_start is None or not str(raw_start).strip()
                amount_blank = raw_amount is None or not str(raw_amount).strip()

                if start_blank and amount_blank:
                    continue
                if start_blank or amount_blank:
                    raise HouseDataImportError(
                        f"Лист «{TARIFF_SHEET}», строка {row_number}, "
                        f"пара {pair_number}: заполните одновременно "
                        "«Действует с» и тариф."
                    )
                if not house_id and not address:
                    raise HouseDataImportError(
                        f"Лист «{TARIFF_SHEET}», строка {row_number}: не указан дом."
                    )

                try:
                    start = _parse_tariff_start_cell(start_cell)
                    amount = normalize_positive_decimal(raw_amount)
                except ValueError as exc:
                    raise HouseDataImportError(
                        f"Лист «{TARIFF_SHEET}», строка {row_number}, "
                        f"пара {pair_number}: {exc}"
                    ) from exc

                identity = house_id or _normalize_address(address)
                duplicate_key = (identity, start.value.isoformat())
                if duplicate_key in seen_tariffs:
                    raise HouseDataImportError(
                        f"Лист «{TARIFF_SHEET}», строка {row_number}, "
                        f"пара {pair_number}: для этого дома уже есть тариф "
                        "с такой датой начала."
                    )
                seen_tariffs.add(duplicate_key)
                tariff_rows.append(
                    ImportedTariffRow(house_id, address, start, amount, row_number, **identity_fields(tuple(c.value for c in cells), tariff_headers))
                )

        if not area_rows and not tariff_rows:
            warnings.append("В файле нет заполненных площадей или тарифов.")

        from .house_extra_data import parse_extra_sheets
        extra_houses, extra_people = parse_extra_sheets(workbook)
        if extra_houses or extra_people:
            warnings = [text for text in warnings if text != "В файле нет заполненных площадей или тарифов."]
        return HouseDataImportResult(
            areas=tuple(area_rows),
            tariffs=tuple(tariff_rows),
            warnings=tuple(warnings),
            houses=extra_houses, people=extra_people,
        )
    finally:
        workbook.close()


def normalize_positive_decimal(value: object) -> str:
    text = str(value or "").replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        decimal = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"некорректное число: {value!r}.") from exc
    if not decimal.is_finite() or decimal <= 0:
        raise ValueError("значение должно быть конечным числом больше нуля.")
    result = format(decimal, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return result.replace(".", ",")


def _tariff_column_pairs(sheet) -> list[tuple[int, int]]:  # noqa: ANN001
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    while headers and not headers[-1]:
        headers.pop()

    if headers[:2] != ["_house_id", "Адрес"] or len(headers) < 4:
        raise HouseDataImportError(
            f"Лист «{sheet.title}» имеет неизвестную структуру. "
            "Сформируйте новый шаблон в приложении."
        )

    identity_start = next((i for i, header in enumerate(headers) if header in IDENTITY_HEADERS), len(headers))
    tariff_headers = headers[2:identity_start]
    if len(tariff_headers) % 2:
        raise HouseDataImportError(
            f"Лист «{sheet.title}»: после адреса должны идти пары "
            "«Действует с» / «Тариф, руб./м² в месяц»."
        )

    pairs: list[tuple[int, int]] = []
    for offset in range(0, len(tariff_headers), 2):
        if tariff_headers[offset : offset + 2] != [
            "Действует с",
            "Тариф, руб./м² в месяц",
        ]:
            pair_number = offset // 2 + 1
            raise HouseDataImportError(
                f"Лист «{sheet.title}»: неверные заголовки в паре {pair_number}. "
                "Ожидаются «Действует с» и «Тариф, руб./м² в месяц»."
            )
        start_column = offset + 3
        pairs.append((start_column, start_column + 1))
    return pairs


def _parse_tariff_start_cell(cell) -> ParsedPeriodDate:  # noqa: ANN001
    if cell is None:
        raise ValueError("Дата не заполнена.")

    value = cell.value
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        number_format = str(getattr(cell, "number_format", "") or "").lower()
        number_format = number_format.replace("\\", "")
        month_only_format = (
            "m" in number_format
            and "y" in number_format
            and "d" not in number_format
        )
        if month_only_format:
            return ParsedPeriodDate(
                value=date(value.year, value.month, 1),
                precision="month",
            )
        return ParsedPeriodDate(value=value, precision="day")

    return parse_period_date(value, boundary="start")


def _cell(cells: tuple[object, ...], column: int):  # noqa: ANN001
    index = column - 1
    return cells[index] if 0 <= index < len(cells) else None


def _cell_value(cells: tuple[object, ...], column: int) -> object | None:
    cell = _cell(cells, column)
    return cell.value if cell is not None else None


def _style_header(sheet) -> None:  # noqa: ANN001
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def _require_headers(sheet, expected: list[str]) -> None:  # noqa: ANN001
    actual = [str(cell.value or "").strip() for cell in sheet[1]]
    if actual[: len(expected)] != expected:
        raise HouseDataImportError(
            f"Лист «{sheet.title}» имеет неизвестную структуру. "
            "Сформируйте новый шаблон в приложении."
        )


def _at(values: tuple[object, ...], index: int) -> object | None:
    return values[index] if index < len(values) else None


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_address(value: str) -> str:
    return " ".join(value.split()).casefold()


def normalize_area(value):
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"некорректная площадь: {value!r}.") from exc
    if not number.is_finite() or number < 0:
        raise ValueError("площадь должна быть конечным неотрицательным числом.")
    return "0" if number == 0 else normalize_positive_decimal(number)


def resolve_house_row(row, houses):
    """Require all supplied identity evidence to agree; never match row numbers."""
    houses = list(houses)
    by_id = {h.id: h for h in houses}
    address = _normalize_address(row.address)
    identities = {key: str(getattr(row, key, None) or "").strip().casefold()
                  for key in ("house_object_guid", "fias_house_guid", "gis_unique_house_number")}
    identities = {key: value for key, value in identities.items() if value}
    if row.house_id in by_id:
        house = by_id[row.house_id]
        if address and _normalize_address(house.address) != address:
            raise ValueError("служебный идентификатор и адрес указывают на разные дома")
        if any(str(getattr(house, key) or "").casefold() != value for key, value in identities.items()):
            raise ValueError("идентификаторы дома противоречат друг другу")
        return house
    # Portable GIS identifiers take precedence after local IDs have changed.
    if identities:
        matches = [h for h in houses if all(str(getattr(h, key) or "").casefold() == value
                                            for key, value in identities.items())]
        if len(matches) == 1 and (not address or _normalize_address(matches[0].address) == address):
            return matches[0]
        raise ValueError("идентификаторы ГИС не удалось однозначно сопоставить с адресом у выбранной УО")
    matches = [h for h in houses if address and _normalize_address(h.address) == address]
    if len(matches) == 1:
        return matches[0]
    raise ValueError("адрес отсутствует у выбранной УО или встречается несколько раз; выберите дом вручную")


def area_values_for_house(row, house):
    """Blank component cells retain existing values before deriving charge area."""
    fields = {name: getattr(row, name) for name in ("residential_area", "non_residential_area")
              if getattr(row, name) is not None}
    area = row.area
    if not row.charge_explicit:
        effective = [fields.get(name, getattr(house, name)) for name in ("residential_area", "non_residential_area")]
        area = normalize_area(sum((Decimal(normalize_area(v).replace(",", "."))
                                   for v in effective if v is not None and str(v).strip()), Decimal(0)))
    return {"charge_area": area, **fields}
