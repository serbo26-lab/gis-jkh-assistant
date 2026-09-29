"""Parser for GIS ЖКХ «Объекты жилищного фонда» XLSX exports.

The parser is intentionally structure-based. It records TYPE/VERSION when present,
but it does not reject a file only because its version is unfamiliar.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class HousingObjectsImportError(ValueError):
    """Raised when the workbook cannot be interpreted reliably."""


@dataclass(frozen=True, slots=True)
class ImportedOrganization:
    name: str
    ogrn: str
    kpp: str | None = None


@dataclass(frozen=True, slots=True)
class ImportedHouse:
    organization_ogrn: str
    address: str
    house_object_guid: str | None
    gis_house_guid: str | None
    fias_house_guid: str | None
    gis_unique_house_number: str | None
    cadastral_number: str | None
    residential_premises_count: int
    non_residential_premises_count: int


@dataclass(frozen=True, slots=True)
class HousingObjectsImportResult:
    file_type: str | None
    version: str | None
    organizations: tuple[ImportedOrganization, ...]
    houses: tuple[ImportedHouse, ...]
    warnings: tuple[str, ...] = ()


@dataclass(slots=True)
class _HouseAccumulator:
    organization_ogrn: str
    address: str | None = None
    house_object_guid: str | None = None
    gis_house_guid: str | None = None
    fias_house_guid: str | None = None
    gis_unique_house_number: str | None = None
    cadastral_number: str | None = None
    residential_premises: set[str] = field(default_factory=set)
    non_residential_premises: set[str] = field(default_factory=set)


_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "address": ("Адрес ОЖФ", "Адрес объекта жилищного фонда"),
    "gis_house_guid": (
        "HOUSEGUID (Идентификационный код дома в ГИС ЖКХ)",
        "Идентификационный код дома в ГИС ЖКХ",
    ),
    "fias_house_guid": (
        "HOUSEGUID (Глобальный уникальный идентификатор дома по ФИАС)",
        "Глобальный уникальный идентификатор дома по ФИАС",
    ),
    "org_ogrn": ("ОГРН управляющей организации", "ОГРН УО"),
    "org_kpp": ("КПП управляющей организации", "КПП УО"),
    "org_name": (
        "Краткое наименование управляющей организации",
        "Наименование управляющей организации",
    ),
    "residential_number": ("Номер жилого помещения",),
    "non_residential_number": ("Номер нежилого помещения",),
    "room_number": ("Номер комнаты",),
    "gis_unique_house_number": ("Уникальный номер дома",),
    "unique_premise_number": ("Уникальный номер помещения",),
    "unique_room_number": ("Уникальный номер комнаты",),
    "cadastral_number": ("Кадастровый номер", "Кадастровый номер "),
    "house_object_guid": ("Глобальный уникальный идентификатор объекта жилищного фонда",),
}

_REQUIRED_FIELDS = {
    "address",
    "org_ogrn",
    "org_name",
    "residential_number",
    "non_residential_number",
    "unique_premise_number",
    "house_object_guid",
}


def parse_housing_objects_export(path: str | Path) -> HousingObjectsImportResult:
    """Read a GIS ЖКХ housing-object export and aggregate it to house cards."""

    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        file_type, version = _read_conf(workbook)
        worksheet, columns = _find_data_sheet(workbook)

        organizations: dict[str, ImportedOrganization] = {}
        houses: dict[tuple[str, str], _HouseAccumulator] = {}
        warnings: list[str] = []

        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in values):
                continue

            ogrn = _text(_cell(values, columns, "org_ogrn"))
            org_name = _text(_cell(values, columns, "org_name"))
            if not ogrn or not org_name:
                continue

            kpp = _text(_cell(values, columns, "org_kpp"))
            organizations[ogrn] = ImportedOrganization(name=org_name, ogrn=ogrn, kpp=kpp)

            gis_house_guid = _text(_cell(values, columns, "gis_house_guid"))
            unique_house_number = _text(_cell(values, columns, "gis_unique_house_number"))
            raw_address = _text(_cell(values, columns, "address"))

            group_id = gis_house_guid or unique_house_number
            if not group_id:
                warnings.append(
                    f"Пропущена строка без идентификатора дома: {raw_address or 'без адреса'}"
                )
                continue

            key = (ogrn, group_id)
            acc = houses.setdefault(
                key,
                _HouseAccumulator(
                    organization_ogrn=ogrn,
                    gis_house_guid=gis_house_guid,
                    gis_unique_house_number=unique_house_number,
                ),
            )

            premise_id = _text(_cell(values, columns, "unique_premise_number"))
            room_id = _text(_cell(values, columns, "unique_room_number"))
            residential_number = _text(_cell(values, columns, "residential_number"))
            non_residential_number = _text(_cell(values, columns, "non_residential_number"))
            room_number = _text(_cell(values, columns, "room_number"))

            is_house_row = not premise_id and not room_id and not (
                residential_number or non_residential_number or room_number
            )

            if is_house_row:
                acc.address = raw_address or acc.address
                acc.house_object_guid = _text(_cell(values, columns, "house_object_guid"))
                acc.gis_house_guid = gis_house_guid or acc.gis_house_guid
                acc.fias_house_guid = _text(_cell(values, columns, "fias_house_guid"))
                acc.gis_unique_house_number = unique_house_number or acc.gis_unique_house_number
                acc.cadastral_number = _text(_cell(values, columns, "cadastral_number"))
                continue

            # Room rows can repeat the parent premise. Sets prevent double counting.
            if premise_id and residential_number:
                acc.residential_premises.add(premise_id)
            if premise_id and non_residential_number:
                acc.non_residential_premises.add(premise_id)

        imported_houses: list[ImportedHouse] = []
        for acc in houses.values():
            if not acc.address:
                warnings.append(
                    "Не найдена строка самого дома для идентификатора "
                    f"{acc.gis_house_guid or acc.gis_unique_house_number}"
                )
                continue
            imported_houses.append(
                ImportedHouse(
                    organization_ogrn=acc.organization_ogrn,
                    address=acc.address,
                    house_object_guid=acc.house_object_guid,
                    gis_house_guid=acc.gis_house_guid,
                    fias_house_guid=acc.fias_house_guid,
                    gis_unique_house_number=acc.gis_unique_house_number,
                    cadastral_number=acc.cadastral_number,
                    residential_premises_count=len(acc.residential_premises),
                    non_residential_premises_count=len(acc.non_residential_premises),
                )
            )

        imported_houses.sort(key=lambda item: (item.organization_ogrn, item.address.casefold()))
        return HousingObjectsImportResult(
            file_type=file_type,
            version=version,
            organizations=tuple(
                sorted(organizations.values(), key=lambda item: item.name.casefold())
            ),
            houses=tuple(imported_houses),
            warnings=tuple(dict.fromkeys(warnings)),
        )
    finally:
        workbook.close()


def _find_data_sheet(workbook: Any) -> tuple[Any, dict[str, int]]:
    best: tuple[Any, dict[str, int]] | None = None
    best_score = -1

    for worksheet in workbook.worksheets:
        if worksheet.max_row < 1:
            continue
        headers = [_normalize_header(cell.value) for cell in worksheet[1]]
        columns = _resolve_columns(headers)
        score = len(_REQUIRED_FIELDS.intersection(columns))
        if score > best_score:
            best = (worksheet, columns)
            best_score = score

    if best is None:
        raise HousingObjectsImportError("В книге не найден лист с данными.")

    worksheet, columns = best
    missing = sorted(_REQUIRED_FIELDS - columns.keys())
    if missing:
        readable = ", ".join(missing)
        raise HousingObjectsImportError(
            "Формат выгрузки не распознан надежно. Не найдены обязательные поля: " + readable
        )
    if "gis_house_guid" not in columns and "gis_unique_house_number" not in columns:
        raise HousingObjectsImportError(
            "Формат выгрузки не распознан надежно: не найден устойчивый идентификатор дома."
        )
    return worksheet, columns


def _resolve_columns(headers: list[str]) -> dict[str, int]:
    result: dict[str, int] = {}
    for semantic, aliases in _HEADER_ALIASES.items():
        normalized_aliases = {_normalize_header(alias) for alias in aliases}
        for index, header in enumerate(headers):
            if header in normalized_aliases:
                result[semantic] = index
                break
    return result


def _read_conf(workbook: Any) -> tuple[str | None, str | None]:
    if "conf" not in workbook.sheetnames:
        return None, None
    sheet = workbook["conf"]
    values: dict[str, str] = {}
    for row in sheet.iter_rows(values_only=True):
        if len(row) < 2:
            continue
        key = _text(row[0])
        value = _text(row[1])
        if key and value:
            values[key.upper()] = value
    return values.get("TYPE"), values.get("VERSION")


def _cell(values: tuple[Any, ...], columns: dict[str, int], semantic: str) -> Any:
    index = columns.get(semantic)
    if index is None or index >= len(values):
        return None
    return values[index]


def _normalize_header(value: Any) -> str:
    text = _text(value) or ""
    return " ".join(text.split()).casefold()


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
